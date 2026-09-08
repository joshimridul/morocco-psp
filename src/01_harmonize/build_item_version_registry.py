#!/usr/bin/env python3
"""Create a Year 3 item-version/form registry from the approved item maps.

The analytic registry omits item text and retains hashes. A protected companion
manifest under work_root/private_manifests retains the Ministry-facing summary.
Neither file is written to GitHub.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "00_inventory"))
from build_y3_registry import read_simple_paths  # noqa: E402
from path_guard import assert_output_path, assert_source_path  # noqa: E402


SUBJECT_SHEETS = ("Arabic", "French", "Maths")
FIELD_ALIASES = {
    "item_id": "Items ID",
    "question_number": "The question number endline",
    "question_summary": "The question text baseline",
    "subject": "The subject",
    "content_domain": "The content domain",
    "cognitive_domain": "The cognitive domain (updated)",
    "grade_origin": "The curricular grade level the question\u00a0is mapped to (Grade of origin)",
    "form_label": "Baseline question paper (by grades)",
    "map_anchor_flag": "anchor",
    "scto_question_id": "SCTO question id",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalized_text(value: Any) -> str:
    text = clean(value).casefold()
    return re.sub(r"\s+", " ", text)


def digest(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def read_sheet_rows(path: Path, sheet: str) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    iterator = worksheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    rows: list[dict[str, str]] = []
    for values in iterator:
        record = {header: clean(value) for header, value in zip(headers, values) if header}
        if any(record.values()):
            rows.append(record)
    workbook.close()
    return rows


def parse_form_grades(scto_question_id: str, form_label: str) -> tuple[list[str], str]:
    match = re.search(r"(?:^|_)N([1-9]+)(?:_|$)", scto_question_id, flags=re.IGNORECASE)
    if match:
        return list(match.group(1)), "scto_question_id"
    match = re.fullmatch(r"Grade\s*([1-9])", form_label, flags=re.IGNORECASE)
    if match:
        return [match.group(1)], "form_label"
    grades = re.findall(r"[1-9]", form_label)
    if grades:
        return sorted(set(grades)), "form_label_multi_grade"
    return [""], "unresolved"


def read_map(path: Path, wave: str) -> tuple[list[dict[str, str]], dict[str, dict[str, str]]]:
    all_rows = read_sheet_rows(path, "All")
    subject_metadata: dict[str, dict[str, str]] = {}
    for sheet in SUBJECT_SHEETS:
        for row in read_sheet_rows(path, sheet):
            key = clean(row.get(FIELD_ALIASES["scto_question_id"], ""))
            if key:
                subject_metadata[key] = row

    result: list[dict[str, str]] = []
    for row in all_rows:
        item_id = clean(row.get(FIELD_ALIASES["item_id"], ""))
        scto_id = clean(row.get(FIELD_ALIASES["scto_question_id"], ""))
        if not item_id or not scto_id:
            continue
        subject_row = subject_metadata.get(scto_id, {})
        merged = {
            key: clean(row.get(header, "")) or clean(subject_row.get(header, ""))
            for key, header in FIELD_ALIASES.items()
        }
        if merged["subject"]:
            merged["subject_source"] = "map"
        else:
            merged["subject"] = {"a": "Arabic", "f": "French", "m": "Maths"}.get(
                item_id[:1].lower(), ""
            )
            merged["subject_source"] = "inferred_from_item_prefix_requires_confirmation"
        form_grades, grade_source = parse_form_grades(merged["scto_question_id"], merged["form_label"])
        prompt_hash = digest(normalized_text(merged["question_summary"])) if merged["question_summary"] else ""
        fingerprint_payload = "|".join(
            [merged["subject"], item_id, prompt_hash, merged["grade_origin"]]
        )
        for form_grade in form_grades:
            expanded = dict(merged)
            expanded.update(
                {
                    "wave": wave,
                    "form_grade": form_grade,
                    "form_grade_source": grade_source,
                    "form_group": "".join(form_grades),
                    "prompt_summary_sha256": prompt_hash,
                    "provisional_version_fingerprint": digest(fingerprint_payload),
                    "fingerprint_scope": "item_id+subject+prompt_summary+grade_origin; options/key/rubric/stimulus/layout pending",
                    "human_version_review_required": "1",
                    "source_map": path.name,
                }
            )
            result.append(expanded)
    return result, subject_metadata


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--github-root", type=Path, required=True)
    args = parser.parse_args()
    paths = read_simple_paths(args.config.resolve())
    source_roots = [paths["y1_root"], paths["y2_root"], paths["y3_root"]]
    y3 = paths["y3_root"]
    github = args.github_root.resolve()
    work = paths["work_root"]

    maps = {
        "baseline": y3 / "3 - Data collection/01_Baseline/3 - Item map/Item_map_20251013_eo.xlsx",
        "pilot": y3 / "3 - Data collection/02_Endline_Pilot/3 - Item map/Item_map_20260523_mji.xlsx",
        "endline": y3 / "3 - Data collection/03_Endline/3 - Item map/Item_map_20260618_eo.xlsx",
    }
    ministry_path = github / "exchange-with-ministry/received-from-ministry/Endline item map_from_btihaj.xlsx"
    for path in maps.values():
        assert_source_path(path, source_roots)
    if not ministry_path.exists():
        raise FileNotFoundError(ministry_path)

    all_rows: list[dict[str, str]] = []
    wave_meta: dict[str, Any] = {}
    for wave, path in maps.items():
        rows, subject_meta = read_map(path, wave)
        all_rows.extend(rows)
        wave_meta[wave] = {
            "rows": len(rows),
            "unique_item_ids": len({row["item_id"] for row in rows}),
            "unique_scto_ids": len({row["scto_question_id"] for row in rows}),
            "subject_sheet_rows_keyed": len(subject_meta),
        }

    ministry_rows, ministry_subject_meta = read_map(ministry_path, "ministry_return")
    ministry_by_scto_grade = {
        (row["scto_question_id"], row["form_grade"]): row for row in ministry_rows
    }
    for row in all_rows:
        if row["wave"] == "endline":
            ministry = ministry_by_scto_grade.get((row["scto_question_id"], row["form_grade"]), {})
            row["ministry_content_domain"] = ministry.get("content_domain", "")
            row["ministry_cognitive_domain"] = ministry.get("cognitive_domain", "")
            row["domain_source"] = "ministry_return_subject_sheet" if (
                row["ministry_content_domain"] or row["ministry_cognitive_domain"]
            ) else "missing"
        else:
            row["ministry_content_domain"] = ""
            row["ministry_cognitive_domain"] = ""
            row["domain_source"] = "wave_map_subject_sheet" if (
                row["content_domain"] or row["cognitive_domain"]
            ) else "missing"

    by_wave_item: dict[tuple[str, str], set[str]] = defaultdict(set)
    waves_by_item: dict[str, set[str]] = defaultdict(set)
    hashes_by_item: dict[str, set[str]] = defaultdict(set)
    for row in all_rows:
        by_wave_item[(row["wave"], row["item_id"])].add(row["form_grade"])
        waves_by_item[row["item_id"]].add(row["wave"])
        if row["prompt_summary_sha256"]:
            hashes_by_item[row["item_id"]].add(row["prompt_summary_sha256"])
    for row in all_rows:
        grades = sorted(grade for grade in by_wave_item[(row["wave"], row["item_id"])] if grade)
        row["administered_grades_for_item_in_wave"] = ",".join(grades)
        row["within_wave_adjacent_grade_candidate"] = "1" if len(grades) >= 2 else "0"
        across = waves_by_item[row["item_id"]]
        row["cross_wave_id_match_candidate"] = "1" if len(across) >= 2 else "0"
        row["cross_wave_prompt_hash_consistent"] = "1" if (
            len(across) >= 2 and len(hashes_by_item[row["item_id"]]) == 1
        ) else "0"
        row["anchor_eligibility_status"] = "requires empirical and full-version review"

    private_fields = [
        "wave", "subject", "subject_source", "form_grade", "form_grade_source", "form_group", "item_id", "scto_question_id",
        "question_number", "question_summary", "grade_origin", "form_label", "map_anchor_flag",
        "content_domain", "cognitive_domain", "ministry_content_domain", "ministry_cognitive_domain",
        "domain_source", "prompt_summary_sha256", "provisional_version_fingerprint", "fingerprint_scope",
        "human_version_review_required", "administered_grades_for_item_in_wave",
        "within_wave_adjacent_grade_candidate", "cross_wave_id_match_candidate",
        "cross_wave_prompt_hash_consistent", "anchor_eligibility_status", "source_map",
    ]
    analytic_fields = [field for field in private_fields if field != "question_summary"]
    private_path = assert_output_path(
        work / "private_manifests/y3_item_version_registry_private.csv", work, source_roots
    )
    analytic_path = assert_output_path(
        work / "derived/y3_ministry/y3_item_version_registry.csv", work, source_roots
    )
    summary_path = assert_output_path(
        work / "outputs/y3_ministry/00_readiness/y3_item_map_integrity.json", work, source_roots
    )
    write_csv(private_path, all_rows, private_fields)
    write_csv(analytic_path, all_rows, analytic_fields)

    endline_rows = [row for row in all_rows if row["wave"] == "endline"]
    summary = {
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "wave_maps": wave_meta,
        "total_form_rows": len(all_rows),
        "ministry_subject_sheet_rows_keyed": len(ministry_subject_meta),
        "endline_rows_with_both_ministry_domains": sum(
            bool(row["ministry_content_domain"] and row["ministry_cognitive_domain"])
            for row in endline_rows
        ),
        "endline_rows_missing_any_ministry_domain": sum(
            not (row["ministry_content_domain"] and row["ministry_cognitive_domain"])
            for row in endline_rows
        ),
        "endline_grade_source_counts": Counter(row["form_grade_source"] for row in endline_rows),
        "endline_form_grade_counts": Counter(row["form_grade"] for row in endline_rows),
        "endline_map_anchor_rows": sum(row["map_anchor_flag"] == "1" for row in endline_rows),
        "endline_within_wave_shared_item_rows": sum(
            row["within_wave_adjacent_grade_candidate"] == "1" for row in endline_rows
        ),
        "cross_wave_id_match_item_count": len(
            {row["item_id"] for row in all_rows if row["cross_wave_id_match_candidate"] == "1"}
        ),
        "cross_wave_prompt_consistent_item_count": len(
            {row["item_id"] for row in all_rows if row["cross_wave_prompt_hash_consistent"] == "1"}
        ),
        "fingerprint_limit": "Prompt-summary hash is not sufficient evidence of an anchor; options, key/rubric, stimulus, layout/context, administration, scoring, and exposure remain to verify.",
    }
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False, default=dict), encoding="utf-8")
    print(json.dumps({"analytic_registry": str(analytic_path), "private_registry": str(private_path), **summary}, ensure_ascii=False, default=dict))


if __name__ == "__main__":
    main()
