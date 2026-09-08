#!/usr/bin/env python3
"""Build provisional Y1/Y2/Y3 trend-anchor candidates from the compiled map.

This is an ID/text-summary/form graph, not proof of a common scale. Full item
version, scoring, exposure, and Y1/Y2 empirical verification remain mandatory.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "00_inventory"))
from build_y3_registry import read_simple_paths  # noqa: E402
from path_guard import assert_output_path, assert_source_path  # noqa: E402


def clean(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalized_hash(value: Any) -> str:
    text = re.sub(r"\s+", " ", clean(value).casefold())
    return hashlib.sha256(text.encode("utf-8")).hexdigest() if text else ""


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def number(value: str) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, required=True)
    args = parser.parse_args()
    paths = read_simple_paths(args.config.resolve())
    source_roots = [paths["y1_root"], paths["y2_root"], paths["y3_root"]]
    work = paths["work_root"]
    compiled = paths["y3_root"] / "3 - Data collection/01_Baseline/3 - Item map/Y1-3 Item Map/Item_maps_compiled_16122025_eo.xlsx"
    assert_source_path(compiled, source_roots)

    workbook = load_workbook(compiled, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    raw_rows = [dict(zip(headers, row)) for row in iterator]
    workbook.close()

    years_by_item: dict[str, set[str]] = defaultdict(set)
    for row in raw_rows:
        item = clean(row.get("LatestItemsID"))
        year = clean(row.get("Year"))
        if item and year:
            years_by_item[item].add(year)

    registry = read_csv_rows(work / "derived/y3_ministry/y3_item_version_registry.csv")
    baseline_registry = {
        (row["item_id"], row["subject"], row["form_grade"]): row
        for row in registry
        if row["wave"] == "baseline"
    }
    evidence = read_csv_rows(work / "derived/y3_ministry/y3_item_decision_evidence.csv")
    baseline_evidence = {
        (row["item_id"], row["subject"], row["grade"]): row
        for row in evidence
        if row["wave"] == "baseline"
    }

    output: list[dict[str, Any]] = []
    for row in raw_rows:
        if clean(row.get("Year")) != "Y3":
            continue
        item = clean(row.get("LatestItemsID"))
        subject = clean(row.get("Subject"))
        grades = sorted(set(re.findall(r"[1-6]", clean(row.get("Grade")))))
        for grade in grades:
            current = baseline_registry.get((item, subject, grade), {})
            empirical = baseline_evidence.get((item, subject, grade), {})
            prior_years = sorted(year for year in years_by_item[item] if year != "Y3")
            compiled_hash = normalized_hash(row.get("Thequestiontext"))
            current_hash = current.get("prompt_summary_sha256", "")
            hash_consistent = bool(compiled_hash and current_hash and compiled_hash == current_hash)
            p_correct = number(empirical.get("pct_correct", ""))
            point_biserial = number(empirical.get("point_biserial_rest", ""))
            negative = empirical.get("negative_discrimination_warning") == "1"
            severe = empirical.get("severe_facility_warning") == "1"
            item_fit = empirical.get("irt_item_fit_warning") == "1"
            all_three_years = {"Y1", "Y2"}.issubset(set(prior_years))

            if not prior_years:
                tier = "not a cross-year ID candidate"
            elif not current:
                tier = "requires mapping reconciliation"
            elif not hash_consistent:
                tier = "requires prompt-summary reconciliation"
            elif p_correct is None or point_biserial is None:
                tier = "requires Year 3 empirical review"
            elif negative or severe or point_biserial < 0.15 or p_correct < 5 or p_correct > 95:
                tier = "not eligible empirically in Year 3"
            elif all_three_years and point_biserial >= 0.25 and 10 <= p_correct <= 90 and not item_fit:
                tier = "Tier A provisional trend candidate"
            else:
                tier = "Tier B provisional trend candidate"

            output.append(
                {
                    "item_id": item,
                    "subject": subject,
                    "administered_grade_y3": grade,
                    "used_in_y3": clean(row.get("Used_in")),
                    "baseline_endline_y3": clean(row.get("BaselineEndline")),
                    "prior_years_with_same_latest_id": ",".join(prior_years),
                    "appears_y1_y2_y3": int(all_three_years),
                    "compiled_content_domain": clean(row.get("Thecontentdomain")),
                    "compiled_cognitive_domain": clean(row.get("Thecognitivedomain")),
                    "compiled_prompt_summary_sha256": compiled_hash,
                    "current_prompt_summary_sha256": current_hash,
                    "prompt_summary_hash_consistent": int(hash_consistent),
                    "year3_pct_correct": empirical.get("pct_correct", ""),
                    "year3_point_biserial_rest": empirical.get("point_biserial_rest", ""),
                    "year3_empirical_quality_status": empirical.get("empirical_quality_status", ""),
                    "year3_local_dependence_warning": empirical.get("local_dependence_item_warning", ""),
                    "year3_item_fit_warning": empirical.get("irt_item_fit_warning", ""),
                    "trend_candidate_tier": tier,
                    "cross_year_empirical_status": "Y1/Y2 response-data validation pending",
                    "full_version_review_status": "prompt/stimulus/options/key-rubric/layout/scoring/exposure review required",
                    "common_scale_claim_status": "not established by compiled map",
                }
            )

    fields = list(output[0].keys())
    out_path = assert_output_path(
        work / "outputs/y3_ministry/02_anchors/y3_across_year_trend_candidates.csv",
        work,
        source_roots,
    )
    write_csv(out_path, output, fields)
    summary = {
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "y3_form_item_rows": len(output),
        "unique_y3_item_ids": len({row["item_id"] for row in output}),
        "unique_item_ids_with_prior_year": len(
            {row["item_id"] for row in output if row["prior_years_with_same_latest_id"]}
        ),
        "unique_item_ids_appearing_y1_y2_y3": len(
            {row["item_id"] for row in output if row["appears_y1_y2_y3"] == 1}
        ),
        "tier_counts": Counter(row["trend_candidate_tier"] for row in output),
        "interpretation": "These are provisional trend-anchor candidates based on the compiled latest-item ID, prompt-summary hash, and Year 3 baseline behavior. They are not evidence of a common scale and require Y1/Y2 response and full-version validation.",
    }
    summary_path = assert_output_path(
        work / "outputs/y3_ministry/02_anchors/y3_across_year_trend_summary.json",
        work,
        source_roots,
    )
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False, default=dict), encoding="utf-8")
    print(json.dumps({"output": str(out_path), **summary}, ensure_ascii=False, default=dict))


if __name__ == "__main__":
    main()
