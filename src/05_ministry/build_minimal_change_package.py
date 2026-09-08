#!/usr/bin/env python3
"""Build a minimal-change anchor, replacement, and tail-coverage package.

The historical item bank remains in read-only Dropbox source roots. Secure item
summaries and candidate evidence are written only to the protected work_root.
The Ministry-facing outputs deliberately keep a small number of decision fields.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "00_inventory"))
from build_y3_registry import read_simple_paths  # noqa: E402
from path_guard import assert_output_path, assert_source_path  # noqa: E402


MIN_ANCHORS_PER_LINK = 5
MIN_EASY_ITEMS_PER_FORM = 2
MIN_HARD_ITEMS_PER_FORM = 2


def clean(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalized(value: Any) -> str:
    value = unicodedata.normalize("NFKD", clean(value).casefold())
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def normalized_id(value: Any) -> str:
    text = re.sub(r"[^a-z0-9]+", "", clean(value).casefold())
    text = re.sub(r"y[123](?:nv|v)?[0-9]*$", "", text)
    text = re.sub(r"(?:nv|v)[0-9]*$", "", text)
    return text


def prompt_hash(value: Any) -> str:
    text = re.sub(r"\s+", " ", normalized(value))
    return hashlib.sha256(text.encode("utf-8")).hexdigest() if text else ""


def as_float(value: Any) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def as_binary(value: Any) -> bool:
    if value in (True, 1, "1"):
        return True
    return normalized(value) in {"yes", "oui", "true", "integer"}


def grade_number(value: Any) -> int | None:
    match = re.search(r"[1-6]", clean(value))
    return int(match.group()) if match else None


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_rows(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def first_value(row: dict[str, Any], candidates: Iterable[str]) -> Any:
    normalized_row = {normalized(key): value for key, value in row.items() if key is not None}
    for candidate in candidates:
        value = normalized_row.get(normalized(candidate))
        if value not in (None, ""):
            return value
    return ""


def iter_sheet_rows(path: Path, sheet_name: str, max_rows: int = 2000) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, max_rows), values_only=True)
    headers = [clean(value) for value in next(iterator)]
    blank_streak = 0
    for values in iterator:
        if not any(value not in (None, "") for value in values[:30]):
            blank_streak += 1
            if blank_streak >= 25:
                break
            continue
        blank_streak = 0
        yield dict(zip(headers, values))
    workbook.close()


def load_compiled_crosswalk(path: Path) -> tuple[dict[str, str], dict[str, set[str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    by_scto: dict[str, str] = {}
    years_by_id: dict[str, set[str]] = defaultdict(set)
    for values in iterator:
        row = dict(zip(headers, values))
        item_id = clean(row.get("LatestItemsID"))
        scto = normalized_id(row.get("SCTOquestionid"))
        year = clean(row.get("Year"))
        if item_id and scto:
            by_scto[scto] = item_id
        if item_id and year:
            years_by_id[normalized_id(item_id)].add(year)
    workbook.close()
    return by_scto, years_by_id


def historical_sources(y1_root: Path) -> list[tuple[str, Path, str]]:
    base = y1_root / "3 - Data collection/Item map"
    return [
        ("Arabic", base / "Arabic/Item map arabic 2024-08-26 seo.xlsx", "Arabic baseline"),
        ("Arabic", base / "Arabic/Item map arabic 2024-08-26 seo.xlsx", "Arabic endline-pilot"),
        ("French", base / "French/item map french 2024-09-11 adb.xlsx", "French baseline"),
        ("French", base / "French/item map french 2024-09-11 adb.xlsx", "French endline-pilot"),
        ("Maths", base / "Math/Item map maths 2024-09-11 adb.xlsx", "Math baseline"),
        ("Maths", base / "Math/Item map maths 2024-09-11 adb.xlsx", "Math endline-pilot"),
    ]


def build_historical_bank(
    sources: list[tuple[str, Path, str]],
    by_scto: dict[str, str],
    years_by_id: dict[str, set[str]],
    current_ids: set[str],
    current_prompt_hashes: set[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    candidates: list[dict[str, Any]] = []
    version_observations: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for subject, path, sheet_name in sources:
        for row in iter_sheet_rows(path, sheet_name):
            bank_question_id = clean(first_value(row, ["Question ID", "Question ID (old)", "Question ID endline"]))
            new_item_id = clean(first_value(row, ["New items ID", "New Items ID"]))
            scto = clean(first_value(row, ["Question label in SCTO", "SCTO question id", "var name in baseline clean"]))
            candidate_id = by_scto.get(normalized_id(scto), "") or new_item_id or bank_question_id
            if not candidate_id:
                continue
            question = clean(first_value(row, ["The question text", "The question text Endline", "The question text endline"]))
            correct_answer = clean(first_value(row, [
                "The correct answer (or grading guidelines)",
                "The correct answer",
                "Correct answer",
                "Grading guidelines",
            ]))
            if question:
                version_observations[(subject, normalized_id(candidate_id))].append({
                    "candidate_id": candidate_id,
                    "question_summary_sha256": prompt_hash(question),
                    "correct_answer_or_guideline": correct_answer,
                    "historical_source": f"Y1 {sheet_name}",
                })
            grade = grade_number(first_value(row, [
                "The curricular grade level the question is mapped to",
                "The curricular grade level the question is mapped to (Grade of origin)",
                "Endline question paper (by grades)",
                "Used for baseline",
            ]))
            content = clean(first_value(row, ["The content domain", "The detailed domain"]))
            cognitive = clean(first_value(row, ["The cognitive domain", "The cognitive domain (updated)"]))
            discrimination = as_float(first_value(row, ["a", "IRT discrimination"]))
            difficulty = as_float(first_value(row, ["b", "IRT difficulty"]))
            p_correct = as_float(first_value(row, ["percentage", "Proportion correct"]))
            itemrest = as_float(first_value(row, ["itemrestcor"]))
            if p_correct is not None and p_correct > 1:
                p_correct /= 100
            p_source = "observed proportion correct"
            if p_correct is None and discrimination is not None and difficulty is not None:
                # For maps that retain only 2PL a/b, use the model-implied probability
                # at theta=0 for ranking. It is explicitly marked as derived evidence.
                exponent = max(-40.0, min(40.0, discrimination * difficulty))
                p_correct = 1 / (1 + math.exp(exponent))
                p_source = "2PL-implied at theta=0"
            integer_flag = first_value(row, ["integer question", "Graded answer is an integer?", "Integer"])
            if grade is None or not question or p_correct is None:
                continue
            if as_binary(integer_flag):
                continue
            if p_correct < 0.10 or p_correct > 0.90:
                continue
            if discrimination is not None and discrimination < 0.70 and (itemrest is None or itemrest < 0.20):
                continue
            if discrimination is None and (itemrest is None or itemrest < 0.20):
                continue
            if itemrest is not None and itemrest < 0.15:
                continue
            base_id = normalized_id(candidate_id)
            question_hash = prompt_hash(question)
            if base_id in current_ids or question_hash in current_prompt_hashes:
                continue
            if p_correct >= 0.70:
                tail_role = "easy / lower-tail information"
            elif p_correct <= 0.35:
                tail_role = "hard / upper-tail information"
            else:
                tail_role = "middle"
            years = sorted(years_by_id.get(base_id, {"Y1"}))
            confidence = "high" if itemrest is not None and itemrest >= 0.25 and (discrimination or 0) >= 0.80 else "moderate"
            candidates.append({
                "candidate_id": candidate_id,
                "candidate_base_id": base_id,
                "bank_question_id": bank_question_id,
                "scto_question_id": scto,
                "subject": subject,
                "grade_origin": grade,
                "content_domain": content,
                "cognitive_domain": cognitive,
                "question_summary": question,
                "correct_answer_or_guideline": correct_answer,
                "question_summary_sha256": question_hash,
                "historical_discrimination_a": discrimination,
                "historical_difficulty_b": difficulty,
                "historical_p_correct": p_correct,
                "historical_p_source": p_source,
                "historical_itemrestcor": itemrest,
                "tail_role": tail_role,
                "years_with_latest_id": ",".join(years),
                "historical_source": f"Y1 {sheet_name}",
                "evidence_confidence": confidence,
                "version_review_status": "exact prompt/stimulus/options/key/rubric/layout/scoring/exposure review required",
            })

    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in candidates:
        grouped[(row["subject"], row["candidate_base_id"])].append(row)

    best: dict[tuple[str, str], dict[str, Any]] = {}
    conflicts: list[dict[str, Any]] = []
    conflicted_keys: set[tuple[str, str]] = set()
    for key, records in version_observations.items():
        prompt_versions = {row["question_summary_sha256"] for row in records if row["question_summary_sha256"]}
        answer_versions = {
            normalized(row["correct_answer_or_guideline"])
            for row in records if normalized(row["correct_answer_or_guideline"])
        }
        if len(prompt_versions) > 1 or len(answer_versions) > 1:
            conflicted_keys.add(key)
            conflicts.append({
                "subject": key[0],
                "candidate_id": records[0]["candidate_id"],
                "candidate_base_id": key[1],
                "prompt_version_n": len(prompt_versions),
                "answer_version_n": len(answer_versions),
                "source_contexts": "; ".join(sorted({row["historical_source"] for row in records})),
                "status": "excluded from automatic selection; exact item-version resolution required",
            })
    for key, records in grouped.items():
        if key in conflicted_keys:
            continue
        row = max(
            records,
            key=lambda record: (
                1 if record["historical_itemrestcor"] is not None else 0,
                record["historical_itemrestcor"] or 0,
                min(record["historical_discrimination_a"] or 0, 4),
            ),
        )
        row["candidate_version_fingerprint"] = hashlib.sha256(
            (row["question_summary_sha256"] + "|" + normalized(row["correct_answer_or_guideline"])).encode("utf-8")
        ).hexdigest()
        best[key] = row
    return list(best.values()), conflicts


def domain_similarity(left: str, right: str) -> float:
    a = set(normalized(left).split())
    b = set(normalized(right).split())
    if not a or not b:
        return 0
    return len(a & b) / len(a | b)


def canonical_content_domain(value: str, subject: str) -> str:
    """Map current short codes and historical labels to one broad domain."""
    text = normalized(value)
    compact = text.replace(" ", "")
    if subject in {"Arabic", "French"}:
        if compact in {"po", "productionoral"} or ("production" in text and "oral" in text):
            return "PO"
        if compact in {"pe", "productionecrit"} or ("production" in text and "ecrit" in text):
            return "PE"
        if compact == "lf" or "decodage" in text or "fluidite" in text:
            return "LF"
        if compact == "co" or (("comprehension" in text) and "oral" in text):
            return "CO"
        if compact == "lc" or "lecture" in text or "comprehension" in text:
            return "LC"
    if subject == "Maths":
        if compact in {"gm", "geometricshapesandmeasures"} or "geometric" in text:
            return "GM"
        if compact == "data" or "data display" in text:
            return "DATA"
        if compact == "rp" or "problem" in text:
            return "RP"
        if "algebr" in text:
            return "ALGEBRIC THINKING"
        if compact in {"calcul", "numeracy", "calculnumeracy"} or "calculation" in text or "numeracy" in text:
            return "CALCUL/NUMERACY"
    return clean(value).upper()


def candidate_score(
    row: dict[str, Any],
    subject: str,
    grades: set[int],
    content_domain: str = "",
    desired_tail: str = "",
    preferred_ids: set[str] | None = None,
    max_grade_distance: int = 1,
) -> float:
    if row["subject"] != subject:
        return -10_000
    distance = min(abs(row["grade_origin"] - grade) for grade in grades)
    if distance > max_grade_distance:
        return -10_000
    score = 8 if distance == 0 else 3 if distance == 1 else 1
    if content_domain:
        if canonical_content_domain(row["content_domain"], subject) == canonical_content_domain(content_domain, subject):
            score += 8
        else:
            score += 2 * domain_similarity(row["content_domain"], content_domain)
    if desired_tail and row["tail_role"].startswith(desired_tail):
        score += 4
    elif desired_tail and desired_tail != "middle" and row["tail_role"] == "middle":
        score += 1
    discrimination = row["historical_discrimination_a"] or 0
    score += min(discrimination, 3)
    # Extremely large legacy 2PL estimates are usually instability signals.
    # Keep them in the protected evidence bank, but rank them below otherwise
    # comparable candidates rather than treating the magnitude as strength.
    if discrimination > 6:
        score -= 25
    score += 4 * (row["historical_itemrestcor"] or 0.20)
    score += 0.5 * len(row["years_with_latest_id"].split(","))
    if preferred_ids and row["candidate_base_id"] in preferred_ids:
        score += 8
    return score


def select_candidates(
    bank: list[dict[str, Any]],
    subject: str,
    grades: set[int],
    n: int,
    content_domain: str = "",
    desired_tail: str = "",
    exclude: set[str] | None = None,
    preferred_ids: set[str] | None = None,
    diversify_domains: bool = False,
    require_content_match: bool = False,
    require_tail_match: bool = False,
    max_grade_distance: int = 1,
) -> list[dict[str, Any]]:
    exclude = exclude or set()
    scored = [
        (candidate_score(
            row, subject, grades, content_domain, desired_tail, preferred_ids,
            max_grade_distance=max_grade_distance,
        ), row)
        for row in bank
        if row["candidate_base_id"] not in exclude
        and (
            not require_content_match
            or canonical_content_domain(row["content_domain"], subject)
            == canonical_content_domain(content_domain, subject)
        )
        and (not require_tail_match or row["tail_role"].startswith(desired_tail))
    ]
    scored = [(score, row) for score, row in scored if score > -1_000]
    scored.sort(key=lambda pair: (pair[0], pair[1]["historical_itemrestcor"] or 0, pair[1]["historical_discrimination_a"] or 0), reverse=True)
    selected: list[dict[str, Any]] = []
    seen_prompts: set[str] = set()
    seen_domains: set[str] = set()
    passes = [True, False] if diversify_domains else [False]
    for require_new_domain in passes:
        for _, row in scored:
            domain = normalized(row["content_domain"])
            if row["question_summary_sha256"] in seen_prompts:
                continue
            if require_new_domain and domain in seen_domains:
                continue
            selected.append(row)
            seen_prompts.add(row["question_summary_sha256"])
            seen_domains.add(domain)
            if len(selected) >= n:
                return selected
    return selected


def short_candidate(row: dict[str, Any]) -> str:
    summary = re.sub(r"\s+", " ", row["question_summary"])
    if len(summary) > 105:
        summary = summary[:102].rstrip() + "…"
    return f"{row['candidate_id']} — {summary}"


def history_text(row: dict[str, Any]) -> str:
    symbol = "≈" if row["historical_p_source"] == "2PL-implied at theta=0" else "="
    bits = [f"p{symbol}{row['historical_p_correct']:.2f}"]
    if row["historical_discrimination_a"] is not None:
        bits.append(f"a={row['historical_discrimination_a']:.2f}")
    if row["historical_itemrestcor"] is not None:
        bits.append(f"r={row['historical_itemrestcor']:.2f}")
    bits.append(row["tail_role"].split(" /")[0])
    return "; ".join(bits)


def desired_tail_from_p(value: Any) -> str:
    p = as_float(value)
    if p is None:
        return "middle"
    p = p / 100 if p > 1 else p
    if p >= 0.70:
        return "easy"
    if p <= 0.35:
        return "hard"
    return "middle"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, required=True)
    args = parser.parse_args()
    paths = read_simple_paths(args.config.resolve())
    roots = [paths["y1_root"], paths["y2_root"], paths["y3_root"]]
    work = paths["work_root"]

    compiled = paths["y3_root"] / "3 - Data collection/01_Baseline/3 - Item map/Y1-3 Item Map/Item_maps_compiled_16122025_eo.xlsx"
    assert_source_path(compiled, roots)
    sources = historical_sources(paths["y1_root"])
    for _, path, _ in sources:
        assert_source_path(path, roots)

    decisions = read_rows(work / "outputs/y3_ministry/03_forms/y3_ministry_item_decisions.csv")
    coverage = read_rows(work / "outputs/y3_ministry/02_anchors/y3_anchor_link_coverage.csv")
    anchor_evidence = read_rows(work / "derived/y3_ministry/y3_anchor_candidate_evidence.csv")
    current_ids = {normalized_id(row["item_id"]) for row in decisions}
    current_prompts = {prompt_hash(row["question_summary"]) for row in decisions if row["question_summary"]}
    by_scto, years_by_id = load_compiled_crosswalk(compiled)
    # Keep historically strong records even when the latest ID is already used in
    # one current form. Reusing that exact item at the missing endpoint is often
    # the smallest defensible anchor repair. Context-specific exclusions below
    # prevent duplication within a form.
    bank, bank_conflicts = build_historical_bank(sources, by_scto, years_by_id, set(), set())
    surface_override_path = work / "private_manifests/manual_surface_change_overrides.csv"
    surface_overrides = read_rows(surface_override_path) if surface_override_path.exists() else []

    current_form_ids: dict[tuple[str, int], set[str]] = defaultdict(set)
    for row in decisions:
        current_form_ids[(row["subject"], int(float(row["administered_grade"])))].add(normalized_id(row["item_id"]))
    registry = read_rows(work / "derived/y3_ministry/y3_item_version_registry.csv")
    baseline_form_ids: dict[tuple[str, int], set[str]] = defaultdict(set)
    for row in registry:
        if row["wave"] == "baseline":
            baseline_form_ids[(row["subject"], int(float(row["form_grade"])))].add(normalized_id(row["item_id"]))

    severe_actions = {
        "revise or replace, then re-pilot": "replace before pilot",
        "retain only after psychometric review and re-pilot": "keep only if key/version review passes; otherwise replace",
    }
    replacement_rows: list[dict[str, Any]] = []
    replacement_selected: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    used_replacement_ids: dict[tuple[str, int], set[str]] = defaultdict(set)
    target_decisions = [row for row in decisions if row["provisional_recommendation"] in severe_actions]
    target_decisions.sort(key=lambda row: (row["subject"], int(float(row["administered_grade"])), row["item_id"]))
    for row in target_decisions:
        grade = int(float(row["administered_grade"]))
        key = (row["subject"], grade)
        desired_tail = desired_tail_from_p(row["pct_correct"])
        selected = select_candidates(
            bank,
            row["subject"],
            {grade},
            2,
            content_domain=row["proposed_content_domain"],
            desired_tail=desired_tail,
            exclude=used_replacement_ids[key] | current_form_ids[key],
            require_content_match=True,
            max_grade_distance=2,
        )
        for candidate in selected:
            used_replacement_ids[key].add(candidate["candidate_base_id"])
            replacement_selected[key].append(candidate)
        primary = selected[0] if selected else None
        backup = selected[1] if len(selected) > 1 else None
        issue_bits = []
        if row["empirical_quality_status"]:
            issue_bits.append(row["empirical_quality_status"])
        if row["school_split_warning"] == "1":
            issue_bits.append("school-split instability")
        if row["local_dependence_warning"] == "1":
            issue_bits.append("shared-task dependence")
        issue = "; ".join(dict.fromkeys(issue_bits)) or "psychometric review warning"
        replacement_rows.append({
            "priority": "change now" if row["provisional_recommendation"] == "revise or replace, then re-pilot" else "conditional backup",
            "subject": row["subject"],
            "grade": grade,
            "current_item": row["item_id"],
            "issue": issue,
            "current_evidence": f"p={as_float(row['pct_correct']) or 0:.1f}%; r={as_float(row['point_biserial_rest']) or 0:.2f}",
            "proposed_action": severe_actions[row["provisional_recommendation"]],
            "primary_candidate": short_candidate(primary) if primary else "no defensible historical match found",
            "primary_history": history_text(primary) if primary else "develop or inspect additional bank sources",
            "backup_candidate": short_candidate(backup) if backup else "",
            "ministry_decision": "",
            "ministry_comment": "",
        })

    current_by_link: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in anchor_evidence:
        if row["anchor_candidate_tier"] in {"Tier A core candidate", "Tier B expanded candidate"}:
            current_by_link[row["link_id"]].append(row)
    tier_order = {"Tier A core candidate": 0, "Tier B expanded candidate": 1}
    for rows in current_by_link.values():
        rows.sort(key=lambda row: (tier_order[row["anchor_candidate_tier"]], row["item_id"]))

    preferred_ids = {candidate["candidate_base_id"] for rows in replacement_selected.values() for candidate in rows}
    anchor_rows: list[dict[str, Any]] = []
    proposed_placements: dict[tuple[str, int], dict[str, dict[str, Any]]] = defaultdict(dict)
    anchor_form_placements_added = 0
    deficient = [row for row in coverage if not row["preliminary_coverage_status"].startswith("at least 5")]
    deficient.sort(key=lambda row: (row["link_type"], row["subject"], int(row["grade_a"]), int(row["grade_b"])))
    for row in deficient:
        grade_a = int(row["grade_a"])
        grade_b = int(row["grade_b"])
        existing = current_by_link.get(row["link_id"], [])
        existing_ids = []
        for candidate in existing:
            if candidate["item_id"] not in existing_ids:
                existing_ids.append(candidate["item_id"])
        needed = max(0, MIN_ANCHORS_PER_LINK - len(existing_ids))
        if row["link_type"] == "within_grade_pre_post":
            endpoint_a = baseline_form_ids[(row["subject"], grade_a)]
            endpoint_b = current_form_ids[(row["subject"], grade_b)]
        else:
            endpoint_a = current_form_ids[(row["subject"], grade_a)]
            endpoint_b = current_form_ids[(row["subject"], grade_b)]
        current_common = endpoint_a & endpoint_b
        one_endpoint_only = endpoint_a ^ endpoint_b
        selected = select_candidates(
            bank,
            row["subject"],
            {grade_a, grade_b},
            needed,
            exclude={normalized_id(value) for value in existing_ids} | current_common,
            preferred_ids=preferred_ids | one_endpoint_only,
            diversify_domains=True,
        )
        for candidate in selected:
            if row["link_type"] == "within_grade_pre_post":
                anchor_form_placements_added += int(candidate["candidate_base_id"] not in endpoint_a)
                anchor_form_placements_added += int(candidate["candidate_base_id"] not in endpoint_b)
                if candidate["candidate_base_id"] not in endpoint_b:
                    proposed_placements[(row["subject"], grade_b)][candidate["candidate_base_id"]] = candidate
            else:
                for grade in {grade_a, grade_b}:
                    if candidate["candidate_base_id"] not in current_form_ids[(row["subject"], grade)]:
                        anchor_form_placements_added += 1
                        proposed_placements[(row["subject"], grade)][candidate["candidate_base_id"]] = candidate
        link_label = (
            f"{row['subject']} G{grade_a} pre/post"
            if row["link_type"] == "within_grade_pre_post"
            else f"{row['subject']} G{grade_a}–G{grade_b}"
        )
        if row["link_type"] == "within_grade_pre_post":
            cycle_note = "future-cycle repair only; already-collected Year 3 waves cannot be retroactively linked"
        else:
            cycle_note = "copy the unchanged candidate only into missing endpoint form(s); freeze the same version in both"
        anchor_rows.append({
            "priority": "critical" if row["preliminary_coverage_status"].startswith("no ") else "targeted repair",
            "link": link_label,
            "current_status": row["preliminary_coverage_status"],
            "retain_current_candidates": ", ".join(existing_ids) if existing_ids else "none",
            "add_from_historical_bank": ", ".join(candidate["candidate_id"] for candidate in selected) if selected else "candidate search unresolved",
            "items_added": len(selected),
            "candidate_total_after_proposal": len(existing_ids) + len(selected),
            "instruction": cycle_note,
            "ministry_decision": "",
            "ministry_comment": "",
        })

    acceptable_current = {
        "retain as scored non-anchor",
        "retain as scored non-anchor; preserve task/testlet bundle",
        "provisional anchor—freeze unchanged",
        "reserve anchor—freeze unchanged pending review",
        # The tail check below assumes the Ministry accepts the proposed
        # non-anchor surface/scoring revisions and those items remain in form.
        "revise administration/wording/scoring, then re-pilot",
    }
    tail_rows: list[dict[str, Any]] = []
    for subject in ["Arabic", "French", "Maths"]:
        for grade in range(1, 7):
            current = [
                row for row in decisions
                if row["subject"] == subject
                and int(float(row["administered_grade"])) == grade
                and row["item_type"] == "binary"
                and row["provisional_recommendation"] in acceptable_current
                and as_float(row["point_biserial_rest"]) is not None
                and as_float(row["point_biserial_rest"]) >= 0.15
            ]
            p_values = [(as_float(row["pct_correct"]) or 0) / 100 for row in current]
            easy_n = sum(0.70 <= p <= 0.90 for p in p_values)
            hard_n = sum(0.10 <= p <= 0.35 for p in p_values)
            missing_easy = max(0, MIN_EASY_ITEMS_PER_FORM - easy_n)
            missing_hard = max(0, MIN_HARD_ITEMS_PER_FORM - hard_n)
            additions: list[dict[str, Any]] = []
            tail_only_ids: set[str] = set()
            placement = proposed_placements[(subject, grade)]
            for tail, missing in [("easy", missing_easy), ("hard", missing_hard)]:
                already = [candidate for candidate in placement.values() if candidate["tail_role"].startswith(tail)]
                used_for_tail = already[:missing]
                need_more = max(0, missing - len(used_for_tail))
                if need_more:
                    selected = select_candidates(
                        bank,
                        subject,
                        {grade},
                        need_more,
                        desired_tail=tail,
                        exclude=set(placement) | current_form_ids[(subject, grade)],
                        preferred_ids=preferred_ids,
                        require_tail_match=True,
                        max_grade_distance=2,
                    )
                    for candidate in selected:
                        placement[candidate["candidate_base_id"]] = candidate
                        tail_only_ids.add(candidate["candidate_base_id"])
                    used_for_tail.extend(selected)
                additions.extend(used_for_tail)
            relevant_additions = []
            seen = set()
            for candidate in additions:
                if candidate["candidate_base_id"] not in seen:
                    relevant_additions.append(candidate)
                    seen.add(candidate["candidate_base_id"])
            unresolved = missing_easy + missing_hard - len(relevant_additions)
            surface_tail_items = [
                row["item_id"] for row in surface_overrides
                if row.get("subject") == subject
                and row.get("target_grade") == str(grade)
                and row.get("role") == "tail support"
            ][:max(0, unresolved)]
            unresolved = max(0, unresolved - len(surface_tail_items))
            if missing_easy == 0 and missing_hard == 0:
                status = "adequate after accepted revisions"
            elif unresolved == 0 and surface_tail_items:
                status = "adequate after proposed surface edit"
            elif unresolved <= 0:
                status = "adequate after proposed bank item(s)"
            else:
                status = "Ministry item adaptation required"
            tail_rows.append({
                "subject": subject,
                "grade": grade,
                "good_easy_items_current": easy_n,
                "good_hard_items_current": hard_n,
                "minimum_each_tail": MIN_EASY_ITEMS_PER_FORM,
                "status": status,
                "proposed_items_for_tail_coverage": ", ".join(candidate["candidate_id"] for candidate in relevant_additions),
                "surface_adjusted_items": ", ".join(surface_tail_items),
                "net_new_tail_only_items": len(tail_only_ids),
                "note": (
                    "easy items inform the lower tail; hard items inform the upper tail; "
                    + (f"{unresolved} required tail role(s) remain unresolved" if unresolved > 0 else "strict tail role satisfied")
                ),
            })

    private_fields = [
        "candidate_id", "bank_question_id", "scto_question_id", "subject", "grade_origin",
        "content_domain", "cognitive_domain", "question_summary", "correct_answer_or_guideline", "historical_discrimination_a",
        "historical_difficulty_b", "historical_p_correct", "historical_p_source", "historical_itemrestcor", "tail_role",
        "years_with_latest_id", "historical_source", "evidence_confidence", "candidate_version_fingerprint", "version_review_status",
    ]
    replacement_fields = [
        "priority", "subject", "grade", "current_item", "issue", "current_evidence", "proposed_action",
        "primary_candidate", "primary_history", "backup_candidate", "ministry_decision", "ministry_comment",
    ]
    anchor_fields = [
        "priority", "link", "current_status", "retain_current_candidates", "add_from_historical_bank",
        "items_added", "candidate_total_after_proposal", "instruction", "ministry_decision", "ministry_comment",
    ]
    tail_fields = [
        "subject", "grade", "good_easy_items_current", "good_hard_items_current", "minimum_each_tail",
        "status", "proposed_items_for_tail_coverage", "surface_adjusted_items", "net_new_tail_only_items", "note",
    ]

    private_path = assert_output_path(work / "private_manifests/historical_bank_candidate_evidence.csv", work, roots)
    conflict_path = assert_output_path(work / "private_manifests/historical_bank_version_conflicts.csv", work, roots)
    replacement_path = assert_output_path(work / "outputs/y3_ministry/03_forms/y3_minimal_replacement_actions.csv", work, roots)
    anchor_path = assert_output_path(work / "outputs/y3_ministry/03_forms/y3_minimal_anchor_repairs.csv", work, roots)
    tail_path = assert_output_path(work / "outputs/y3_ministry/03_forms/y3_tail_coverage_check.csv", work, roots)
    write_rows(private_path, sorted(bank, key=lambda row: (row["subject"], row["grade_origin"], row["candidate_id"])), private_fields)
    write_rows(conflict_path, sorted(bank_conflicts, key=lambda row: (row["subject"], row["candidate_id"])), [
        "subject", "candidate_id", "candidate_base_id", "prompt_version_n", "answer_version_n",
        "source_contexts", "status",
    ])
    write_rows(replacement_path, replacement_rows, replacement_fields)
    write_rows(anchor_path, anchor_rows, anchor_fields)
    write_rows(tail_path, tail_rows, tail_fields)

    unique_anchor_additions = {
        normalized_id(candidate)
        for row in anchor_rows
        for candidate in row["add_from_historical_bank"].split(", ")
        if candidate and candidate != "candidate search unresolved"
    }
    summary = {
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "historical_candidates_screened_in": len(bank),
        "historical_ids_excluded_for_version_conflict": len(bank_conflicts),
        "mandatory_replacement_rows": sum(row["priority"] == "change now" for row in replacement_rows),
        "conditional_backup_rows": sum(row["priority"] == "conditional backup" for row in replacement_rows),
        "deficient_links": len(anchor_rows),
        "historical_anchor_candidates_added": sum(row["items_added"] for row in anchor_rows),
        "anchor_form_placements_added": anchor_form_placements_added,
        "unique_historical_anchor_items_requested": len(unique_anchor_additions),
        "tail_forms_currently_adequate": sum(row["status"] == "adequate after accepted revisions" for row in tail_rows),
        "tail_forms_requiring_repair": sum(row["status"] != "adequate after accepted revisions" for row in tail_rows),
        "tail_forms_fully_repaired_by_bank": sum(row["status"] == "adequate after proposed bank item(s)" for row in tail_rows),
        "tail_forms_fully_repaired_by_surface_edit": sum(row["status"] == "adequate after proposed surface edit" for row in tail_rows),
        "tail_forms_requiring_ministry_adaptation": sum(row["status"] == "Ministry item adaptation required" for row in tail_rows),
        "net_new_tail_only_items": sum(int(row["net_new_tail_only_items"]) for row in tail_rows),
        "minimal_change_rule": "Replace only the strongest problem items; keep six conditional backups; add only enough link items to reach five candidates; reuse anchor/replacement additions for tail coverage before adding any tail-only item.",
        "decision_boundary": "Historical performance supports a pilot shortlist, not automatic reuse. Exact item-version, scoring, exposure, content, and administration review remains required.",
    }
    summary_path = assert_output_path(work / "outputs/y3_ministry/03_forms/y3_minimal_change_summary.json", work, roots)
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
