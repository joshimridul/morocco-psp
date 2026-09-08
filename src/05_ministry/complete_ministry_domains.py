#!/usr/bin/env python3
"""Complete the Ministry domain table using reproducible, provenance-tagged rules.

No classification is silently invented. Existing Ministry codes take priority,
followed by the same item in another endline grade, exact pilot item/form evidence,
the compiled Y1-Y3 map, and finally a small transparent rule set over the protected
English item summary. Rule-derived values are explicitly marked for confirmation.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "00_inventory"))
from build_y3_registry import read_simple_paths  # noqa: E402
from path_guard import assert_output_path, assert_source_path  # noqa: E402


def clean(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def normalize_content(value: str, subject: str) -> str:
    raw = clean(value)
    key = raw.casefold().replace("’", "'")
    compact = re.sub(r"\s+", " ", key)
    if subject in {"Arabic", "French"}:
        if raw.upper() in {"LF", "PE", "PO", "LC", "CO"}:
            return raw.upper()
        if "fluid" in compact or "décodage" in compact or "decodage" in compact:
            return "LF"
        if "production" in compact and ("oral" in compact or "l'oral" in compact):
            return "PO"
        if "production" in compact and ("écrit" in compact or "ecrit" in compact):
            return "PE"
        if "compréhension de l'oral" in compact or "comprehension de l'oral" in compact:
            return "CO"
        if "lecture" in compact and ("compréhension" in compact or "comprehension" in compact):
            return "LC"
    if subject == "Maths":
        upper = raw.upper()
        if upper in {"CALCUL", "NUMERACY", "GM", "DATA", "RP", "ALGEBRIC THINKING"}:
            return upper
        if compact == "calcul":
            return "CALCUL"
        if "numeracy" in compact and "calculation" not in compact:
            return "NUMERACY"
        if "calculation" in compact and "numeracy" in compact:
            return "CALCUL/NUMERACY"
        if "geometric" in compact or "measure" in compact:
            return "GM"
        if "data" in compact:
            return "DATA"
        if "résolution" in compact or "resolution" in compact or "problem" in compact:
            return "RP"
        if "algebr" in compact:
            return "ALGEBRIC THINKING"
    return raw


def normalize_cognitive(value: str) -> str:
    raw = clean(value)
    key = raw.casefold()
    mapping = {
        "knowledge": "KNOWL", "knowl": "KNOWL",
        "application": "APPL", "appl": "APPL",
        "comprehension": "COMP", "comp": "COMP",
        "analysis": "ANALYSE", "analyse": "ANALYSE",
        "synthesis": "SYNT", "synt": "SYNT",
        "reasoning": "REASON", "reason": "REASON", "reaso": "REASON",
        "evaluation": "EVAL", "eval": "EVAL",
    }
    return mapping.get(key, raw.upper())


def heuristic_from_summary(summary: str, subject: str) -> tuple[str, str, str]:
    text = clean(summary).casefold()
    if subject in {"Arabic", "French"} and any(
        token in text for token in (
            "reading comprehension", "understanding informational text", "from the text", "historical monument"
        )
    ):
        content = "LC"
        if any(token in text for token in ("right order", "imagine", "give him advice", "give advice")):
            cognitive = "SYNT"
        elif any(token in text for token in ("proof", "indicates", "writer mean", "main idea", "differences", "why")):
            cognitive = "ANALYSE"
        else:
            cognitive = "COMP"
        return content, cognitive, "rule-derived from protected item summary; requires collaborator confirmation"
    return "", "", "no defensible generic rule"


def compiled_rows(path: Path) -> dict[tuple[str, str, str], list[tuple[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    result: dict[tuple[str, str, str], list[tuple[str, str]]] = defaultdict(list)
    for values in iterator:
        row = dict(zip(headers, values))
        if clean(row.get("Year")) != "Y3":
            continue
        item = clean(row.get("LatestItemsID"))
        subject = clean(row.get("Subject"))
        for grade in sorted(set(re.findall(r"[1-6]", clean(row.get("Grade"))))):
            result[(item, subject, grade)].append(
                (clean(row.get("Thecontentdomain")), clean(row.get("Thecognitivedomain")))
            )
    workbook.close()
    return result


def unique_pair(pairs: list[tuple[str, str]], subject: str) -> tuple[str, str] | None:
    normalized = {
        (normalize_content(content, subject), normalize_cognitive(cognitive))
        for content, cognitive in pairs
        if clean(content) and clean(cognitive)
    }
    return next(iter(normalized)) if len(normalized) == 1 else None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, required=True)
    args = parser.parse_args()
    paths = read_simple_paths(args.config.resolve())
    work = paths["work_root"]
    source_roots = [paths["y1_root"], paths["y2_root"], paths["y3_root"]]
    compiled_path = paths["y3_root"] / "3 - Data collection/01_Baseline/3 - Item map/Y1-3 Item Map/Item_maps_compiled_16122025_eo.xlsx"
    assert_source_path(compiled_path, source_roots)

    rows = read_csv_rows(work / "private_manifests/y3_item_version_registry_private.csv")
    endline = [row for row in rows if row["wave"] == "endline"]
    pilot = [row for row in rows if row["wave"] == "pilot"]
    compiled = compiled_rows(compiled_path)

    same_endline: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    for row in endline:
        if row["ministry_content_domain"] and row["ministry_cognitive_domain"]:
            same_endline[(row["item_id"], row["subject"])].append(
                (row["ministry_content_domain"], row["ministry_cognitive_domain"])
            )
    pilot_exact: dict[tuple[str, str, str], list[tuple[str, str]]] = defaultdict(list)
    for row in pilot:
        pilot_exact[(row["item_id"], row["subject"], row["form_grade"])].append(
            (row["content_domain"], row["cognitive_domain"])
        )

    output: list[dict[str, Any]] = []
    for row in endline:
        content = normalize_content(row["ministry_content_domain"], row["subject"])
        cognitive = normalize_cognitive(row["ministry_cognitive_domain"])
        source = "Ministry return"
        confidence = "Ministry-coded; taxonomy normalized"
        if not (content and cognitive):
            pair = unique_pair(same_endline.get((row["item_id"], row["subject"]), []), row["subject"])
            if pair:
                content, cognitive = pair
                source = "same endline item used in another grade"
                confidence = "high item-ID evidence; form-context confirmation required"
            else:
                pair = unique_pair(
                    pilot_exact.get((row["item_id"], row["subject"], row["form_grade"]), []),
                    row["subject"],
                )
                if pair:
                    content, cognitive = pair
                    source = "exact pilot item + administered grade"
                    confidence = "high metadata evidence; Ministry confirmation required"
                else:
                    pair = unique_pair(
                        compiled.get((row["item_id"], row["subject"], row["form_grade"]), []),
                        row["subject"],
                    )
                    if pair:
                        content, cognitive = pair
                        source = "compiled Y1-Y3 map exact item + grade"
                        confidence = "medium metadata evidence; Ministry confirmation required"
                    else:
                        content, cognitive, confidence = heuristic_from_summary(
                            row["question_summary"], row["subject"]
                        )
                        source = "transparent summary-text rule" if content and cognitive else "unresolved"

        result = dict(row)
        result.update(
            {
                "proposed_content_domain": content,
                "proposed_cognitive_domain": cognitive,
                "domain_completion_source": source,
                "domain_completion_confidence": confidence,
                "requires_ministry_confirmation": int(source != "Ministry return"),
            }
        )
        output.append(result)

    fields = list(output[0].keys())
    out_path = assert_output_path(
        work / "outputs/y3_ministry/03_forms/y3_completed_domain_map.csv", work, source_roots
    )
    write_csv(out_path, output, fields)
    summary = {
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "item_form_grade_rows": len(output),
        "both_domains_complete": sum(
            bool(row["proposed_content_domain"] and row["proposed_cognitive_domain"])
            for row in output
        ),
        "remaining_incomplete": sum(
            not (row["proposed_content_domain"] and row["proposed_cognitive_domain"])
            for row in output
        ),
        "source_counts": Counter(row["domain_completion_source"] for row in output),
        "confirmation_required_rows": sum(row["requires_ministry_confirmation"] == 1 for row in output),
        "normalization_note": "French/Arabic codes use LF, LC, CO, PO, PE; cognitive codes use KNOWL, COMP, APPL, ANALYSE, SYNT, REASON, EVAL. Calculation and Numeracy remains explicitly combined where the source evidence does not support disaggregation.",
    }
    summary_path = assert_output_path(
        work / "outputs/y3_ministry/03_forms/y3_domain_completion_summary.json", work, source_roots
    )
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False, default=dict), encoding="utf-8")
    print(json.dumps({"output": str(out_path), **summary}, ensure_ascii=False, default=dict))


if __name__ == "__main__":
    main()
