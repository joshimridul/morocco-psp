#!/usr/bin/env python3
"""Build the protected Year 3 technical report and Ministry decision memo."""

from __future__ import annotations

import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
BLUE = "2E74B5"
PALE_BLUE = "DDEBF7"
PALE_GOLD = "FFF2CC"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "666666"
WHITE = "FFFFFF"


def load_json(path: Path):
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def load_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_cell_fill(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_width(cell, width_inches: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_document(doc: Document, memo: bool = False):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78 if memo else 0.82)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5 if memo else 9.4)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.06

    for name, size, color, before, after in [
        ("Title", 22, NAVY, 0, 10),
        ("Heading 1", 16, BLUE, 12, 5),
        ("Heading 2", 12.5, NAVY, 9, 4),
        ("Heading 3", 10.5, NAVY, 7, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "MOROCCO PIONEER SCHOOLS  |  YEAR 3 MEASUREMENT AUDIT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("CONFIDENTIAL — protected work output   •   ")
    r.font.name = "Arial"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_field(footer, "PAGE")


def add_masthead(doc: Document, title: str, subtitle: str, label: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(label.upper())
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(5)
    p.add_run(title)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(subtitle)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    rule = doc.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    rule.autofit = False
    set_cell_width(rule.cell(0, 0), 6.9)
    set_cell_fill(rule.cell(0, 0), NAVY)
    rule.cell(0, 0).paragraphs[0].paragraph_format.space_after = Pt(0)
    rule.cell(0, 0).paragraphs[0].add_run(" ").font.size = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, heading: str, text: str, color: str = PALE_GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 6.8)
    set_cell_fill(cell, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(heading + "  ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items, compact: bool = True):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        p.paragraph_format.space_after = Pt(2 if compact else 4)
        p.add_run(item)


def add_numbered(doc: Document, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(3)
        prefix = p.add_run(f"{index}. ")
        prefix.bold = True
        p.add_run(item)


def add_table(doc: Document, headers, rows, widths=None, font_size=8.1, header_color=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for index, header in enumerate(headers):
        cell = hdr.cells[index]
        set_cell_fill(cell, header_color)
        if widths:
            set_cell_width(cell, widths[index])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(header))
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(font_size)
    for ridx, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if widths:
                set_cell_width(cell, widths[index])
            if ridx % 2 == 0:
                set_cell_fill(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_key_value_table(doc: Document, rows):
    add_table(doc, ["Measure", "Result", "Interpretation"], rows, [2.0, 1.1, 3.7], 8.3, BLUE)


def add_source_note(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Evidence note: " + text)
    run.font.size = Pt(7.5)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_page_break(doc: Document):
    doc.add_page_break()


def subject_alpha_rows(form_rows):
    grouped = defaultdict(list)
    for row in form_rows:
        grouped[(row["wave"], row["subject"])].append(float(row["alpha"]))
    return [
        (wave.title(), subject, len(values), f"{min(values):.3f}–{max(values):.3f}")
        for (wave, subject), values in sorted(grouped.items())
    ]


def build_report(work_root: Path, output_path: Path):
    outputs = work_root / "outputs/y3_ministry"
    score = load_json(outputs / "00_readiness/y3_score_reconstruction_summary.json")
    diag = load_json(outputs / "01_item_quality/y3_diagnostics_summary.json")
    uncertainty = load_json(outputs / "01_item_quality/y3_uncertainty_summary.json")
    anchor = load_json(outputs / "02_anchors/y3_anchor_summary.json")
    trend = load_json(outputs / "02_anchors/y3_across_year_trend_summary.json")
    domain = load_json(outputs / "03_forms/y3_domain_completion_summary.json")
    ministry = load_json(outputs / "03_forms/y3_ministry_decision_summary.json")
    minimal = load_json(outputs / "03_forms/y3_minimal_change_summary.json")
    form_rows = load_csv(outputs / "01_item_quality/y3_form_reliability_uncertainty.csv")
    link_rows = load_csv(outputs / "02_anchors/y3_anchor_link_coverage.csv")
    blocking = load_csv(outputs / "00_readiness/y3_blocking_issues.csv")

    alphas = [float(r["alpha"]) for r in form_rows]
    critical_links = [r for r in link_rows if r["preliminary_coverage_status"].startswith("no ")]
    decision_counts = ministry["decision_counts"]

    doc = Document()
    configure_document(doc)
    add_masthead(
        doc,
        "Year 3 Item, Anchor, and Instrument Audit",
        "Technical evidence report • Baseline, pilot, and endline • 4 August 2026",
        "Morocco Pioneer Schools",
    )
    add_callout(
        doc,
        "Bottom line.",
        "The Year 3 forms are generally reliable, but they are not yet ready to be treated as a fully linked measurement system. Item quality, content necessity, and anchor eligibility must remain separate decisions. The next form should freeze only verified link-specific anchors, repair weak or disconnected links, and add explicit testlet/rubric metadata.",
        PALE_BLUE,
    )
    add_key_value_table(doc, [
        ("Forms audited", score["form_count"], "18 subject-grade forms in each of three waves"),
        ("Deduplicated administrations", f"{score['deduplicated_total']:,}", f"{score['duplicate_rows_removed']} excess attempts removed deterministically"),
        ("Binary item×form rows", f"{score['binary_item_form_rows']:,}", "Primary CTT/IRT evidence base"),
        ("Reliability range", f"{min(alphas):.3f}–{max(alphas):.3f}", "Strong internal consistency; not proof of one dimension"),
        ("Required links", anchor["required_link_count"], "16 adequate, 12 weak, 3 unsupported, 2 disconnected"),
        ("Endline domain rows", domain["item_form_grade_rows"], f"All populated; {domain['confirmation_required_rows']} require Ministry confirmation"),
    ])
    doc.add_heading("Decision boundary", level=2)
    add_bullets(doc, [
        "No item was kept or removed solely because of one statistic or p-value.",
        "Treatment assignment was not used in item screening or anchor selection; treatment-related DIF is reserved for a later sensitivity analysis after anchors are frozen.",
        "Candidate anchors remain provisional until prompt, stimulus, options, key/rubric, layout, scoring, administration, and exposure are confirmed.",
        "PII cleanup is team-owned and was excluded from this work. No identifiers were exported.",
    ])

    add_page_break(doc)
    doc.add_heading("1. Scope, source control, and reproducibility", level=1)
    doc.add_paragraph(
        "The audit begins with verified wave-specific Year 3 files, not the legacy combined Y1–Y3 dataset. A SHA-256 source registry records 192 candidates, including 144 retained candidates and 48 deliberately excluded archive, conflicted, recovered, or PII-like files. Fourteen issues remain in the blocking register. Source paths are represented by root aliases and relative paths; all source systems were treated as read-only."
    )
    doc.add_heading("Analytic unit", level=2)
    doc.add_paragraph(
        "The operational unit is an item-version within a form, administered grade, wave, sample/cohort, and scoring context. Shared French forms are expanded to the grades in which they were administered. The audit never collapses evidence by item ID alone."
    )
    doc.add_heading("Scoring and duplicate handling", level=2)
    add_bullets(doc, [
        "Binary responses: 1 = correct, 0 = incorrect; tagged ‘don’t know’ and untagged blank are separately audited but score zero in the primary raw score.",
        "Continuous fluency/rubric fields are reported separately and excluded from binary alpha and IRT.",
        "Within wave × subject × administered grade × panel ID, the retained attempt has the most mapped responses, then longest duration, then earliest source row.",
        "All 1,933 mapped operational item×form rows were reconstructed; 1,906 are binary and 27 are continuous/rubric. No off-form responses were found.",
    ])
    add_source_note(doc, "Independent Stata schema/count validation confirms 15,645 baseline rows, 3,523 pilot rows, and 20,782 endline rows before deterministic deduplication.")
    doc.add_heading("Baseline tie-out", level=2)
    doc.add_paragraph(
        "Among 647 binary baseline item rows, 645 matched prior-output denominators. Of these, 545 reproduced all reported percentages exactly. Remaining differences are small and confined to the documented duplicate-attempt retention difference; the maximum observed percentage-point discrepancy is approximately 0.315."
    )

    add_page_break(doc)
    doc.add_heading("2. Form-level psychometric properties", level=1)
    doc.add_paragraph(
        "Internal consistency is high throughout the assessment system, but its interpretation is conditional on item bundles and dimensionality. School-cluster bootstrap intervals use 200 school resamples; means and proportions use school-cluster sandwich uncertainty."
    )
    add_table(doc, ["Wave", "Subject", "Forms", "Alpha range"], subject_alpha_rows(form_rows), [1.25, 1.75, 1.0, 2.8], 8.4, BLUE)
    doc.add_heading("Dimensionality and model selection", level=2)
    add_bullets(doc, [
        f"Rasch and 2PL models produced {diag['fitted_model_rows']} model rows; {diag['converged_model_rows']} converged.",
        "One-dimensional 2PL is primary at baseline/endline when it converges; Rasch remains the sensitivity model.",
        "Baseline Arabic Grade 1 and French Grade 1 2PL models did not converge, so their primary interpretation uses the Rasch fallback.",
        "Pilot forms use Rasch only because roughly 180–200 respondents per form do not support stable free discrimination estimates for 26–52 items.",
        "Eigenvalue ratios are generally supportive, but some forms are near or below the usual heuristic of 3. High alpha should not be read as proof of essential unidimensionality.",
    ])
    add_callout(doc, "Interpretation.", "The forms can support useful total-score and item-quality evidence now. Linked IRT claims should be limited to links with adequate verified anchors and should be re-estimated after final form assembly.")

    add_page_break(doc)
    doc.add_heading("3. Item-level evidence and local structure", level=1)
    doc.add_paragraph(
        "Item evidence combines operational rates, item-total relationships, primary IRT parameters and fit, nonresponse/don’t-know behavior, and school-split stability. Each threshold is a review flag, not a mechanical rule."
    )
    add_key_value_table(doc, [
        ("Primary parameter rows", f"{diag['primary_item_parameter_rows']:,}", "One row per binary item×form context"),
        ("No empirical warning", f"{diag['empirical_status_counts']['no empirical warning']:,}", "Strongest empirical subset; still needs content/version review"),
        ("One warning", f"{diag['empirical_status_counts']['one warning']:,}", "Retain/revise judgment depends on warning and content role"),
        ("Multiple warnings", f"{diag['empirical_status_counts']['multiple warnings']:,}", "Requires substantive and psychometric review"),
        ("School-split warnings", f"{diag['school_split_warning_rows']} / {diag['school_split_item_rows']}", "Most endline parameters were stable across school halves"),
        ("Q3 warning pairs", f"{diag['local_dependence_warning_pairs']:,}", "Likely shared passages, images, oral tasks, or rubric bundles"),
    ])
    doc.add_heading("Local dependence is the main structural warning", level=2)
    doc.add_paragraph(
        "The 2,079 adjusted-Q3 warning pairs should not be interpreted as 2,079 defective item pairs. The more plausible system-level explanation is undocumented testlet structure: items share stories, pictures, oral prompts, or rubrics. Because the current maps do not reliably provide stimulus_id, testlet_id, or rubric_id, the audit cannot yet separate intended bundling from unintended local dependence."
    )
    add_bullets(doc, [
        "Assign stimulus_id/testlet_id/rubric_id before final form assembly.",
        "Preserve passage/task bundles during item review; rubric criteria are not independent questions.",
        "Re-estimate reliability, dimensionality, and IRT using a testlet/bifactor or simpler bundle-aware fallback where sample size supports it.",
    ])

    add_page_break(doc)
    doc.add_heading("4. Within-year anchor evidence", level=1)
    doc.add_paragraph(
        "Thirty-three links are required: 18 within-grade baseline/endline links and 15 adjacent-grade links. Thirty-one are empirically estimable. Candidate screening is treatment-blind and uses prompt-summary consistency, acceptable endpoint behavior, and school-clustered logistic DIF conditioned on the common-item rest score."
    )
    add_key_value_table(doc, [
        ("Adequate candidate pool", 16, "At least five Tier A/B candidates"),
        ("Weak candidate pool", 12, "Fewer than five Tier A/B candidates"),
        ("No eligible candidate", 3, "Common items exist, but none pass provisional empirical screen"),
        ("Disconnected", 2, "No common binary item exists"),
        ("Tier A link-item rows", 207, "Core candidates; link-specific, not globally reusable"),
        ("Tier B link-item rows", 38, "Expanded candidates for sensitivity or repair"),
    ])
    doc.add_heading("Critical gaps", level=2)
    critical_display = []
    for row in critical_links:
        link_name = row["link_id"].replace("within_grade_pre_post|", "Pre/post: ").replace("adjacent_grade_vertical|", "Adjacent grades: ").replace("|", " ")
        critical_display.append((link_name, row["common_binary_item_n"], row["preliminary_coverage_status"]))
    add_table(doc, ["Link", "Common binary items", "Status"], critical_display, [3.4, 1.2, 2.2], 8.0, BLUE)
    add_bullets(doc, [
        "Arabic Grade 5 pre/post and Arabic Grades 4–5 are disconnected and require deliberate bridge-item insertion.",
        "Mathematics Grade 2 pre/post, Mathematics Grade 6 pre/post, and Mathematics Grades 3–4 have common items but no empirically eligible provisional candidate.",
        f"Retain current eligible anchors and add only enough exact, content-balanced bank items to bring each deficient link to five pilot candidates. The compact repair plan covers all {minimal['deficient_links']} deficient links.",
    ])

    add_page_break(doc)
    doc.add_heading("5. Anchor-set sensitivity and across-year trend", level=1)
    doc.add_heading("Sensitivity", level=2)
    doc.add_paragraph(
        f"The audit estimates {anchor['sensitivity_rows']} core, expanded, and all-eligible mean–sigma sensitivity rows. Sets with fewer than three usable item parameters are intentionally left unestimated. Parameter correlations and link RMSE are reported in the workbook so the team can see where conclusions depend on anchor composition."
    )
    doc.add_heading("Across-year trend candidates", level=2)
    add_key_value_table(doc, [
        ("Unique Year 3 IDs", trend["unique_y3_item_ids"], "Compiled Year 3 baseline map"),
        ("IDs with a prior-year appearance", trend["unique_item_ids_with_prior_year"], "ID-level candidates only"),
        ("IDs appearing Y1/Y2/Y3", trend["unique_item_ids_appearing_y1_y2_y3"], "Potential trend bridges"),
        ("Tier A provisional rows", trend["tier_counts"]["Tier A provisional trend candidate"], "Need Y1/Y2 response and full-version review"),
        ("Tier B provisional rows", trend["tier_counts"]["Tier B provisional trend candidate"], "Sensitivity pool"),
    ])
    add_callout(
        doc,
        "Do not claim a common scale yet.",
        "A matching latest-item ID or prompt-summary hash is not enough. Cross-year anchors require response-data validation plus exact prompt, stimulus, options, key/rubric, layout, scoring, administration, and exposure equivalence. Year 1 serves a different paper and must not be forced onto a common scale.",
    )

    add_page_break(doc)
    doc.add_heading("6. Completing the Ministry’s unfinished item work", level=1)
    doc.add_paragraph(
        "The Ministry return populated both content and cognitive domains for 418 of 519 original rows. The current package expands to the actual administered-grade structure and proposes both domains for all 639 endline item×grade rows. Provenance and confidence are retained row by row."
    )
    add_table(doc, ["Completion source", "Rows", "Required action"], [
        ("Ministry return", 538, "Use subject to blueprint adjudication"),
        ("Exact pilot item + grade", 74, "Confirm inherited classification"),
        ("Same endline item, another grade", 7, "Confirm grade appropriateness"),
        ("Compiled Y1–Y3 map", 5, "Confirm current-version applicability"),
        ("Transparent summary-text rule", 15, "Ministry/content expert must adjudicate"),
    ], [2.5, 0.8, 3.5], 8.2, BLUE)
    add_callout(doc, "Ministry confirmation required.", f"{domain['confirmation_required_rows']} rows are inherited or rule-derived. They are filled so work can proceed, but they are not represented as Ministry-approved classifications.")
    doc.add_heading("Coding normalization", level=2)
    doc.add_paragraph(
        "Language content codes are normalized to LF, LC, CO, PO, and PE. Cognitive codes are normalized to KNOWL, COMP, APPL, ANALYSE, SYNT, REASON, and EVAL. Calculation/Numeracy remains combined when available evidence does not support disaggregation."
    )

    add_page_break(doc)
    doc.add_heading("7. Proposed item dispositions and grade shift", level=1)
    decision_rows = sorted(decision_counts.items(), key=lambda item: (-item[1], item[0]))
    add_table(doc, ["Provisional disposition", "Rows"], [(k, v) for k, v in decision_rows], [5.8, 1.0], 8.2, BLUE)
    doc.add_paragraph(
        "These are endline item×grade dispositions, not Ministry final decisions. ‘Freeze’ means preserve the version unchanged while prompt/stimulus/options/key/rubric/layout/scoring/exposure are checked; it does not mean permanent retention."
    )
    doc.add_heading("One-grade baseline shift", level=2)
    doc.add_paragraph(
        "The proposed baseline shift contains 652 rows: Grade 1 retains the current baseline Grade 1 source, while target baseline Grades 2–6 draw from current endline Grades 1–5. The majority require wording/administration review and a target-grade pilot. Thirty-five items carry an explicit risk of becoming too easy after the shift, and six fluency/rubric fields remain separate components."
    )
    add_callout(doc, "Required before use.", "Approve grade-specific blueprint weights, reconstruct and verify every score/key, assemble the shifted forms, then pilot by school at the target grade. The shift should not be treated as a clerical relabeling exercise.")

    add_page_break(doc)
    doc.add_heading("8. French form redesign", level=1)
    doc.add_paragraph(
        "French currently relies on shared pools for Grades 2–3 and Grades 4–6. Shared pools produce many apparent vertical links, but they also blur grade-specific targeting, content coverage, and exposure. The recommendation is six distinct operational forms with a deliberately controlled bridge subset."
    )
    add_table(doc, ["Grade", "Current evidence pool", "Operational recommendation"], [
        (1, "N1", "Distinct Grade 1 form + controlled link to Grade 2"),
        (2, "N23", "Distinct Grade 2 form + bridges to Grades 1 and 3"),
        (3, "N23", "Distinct Grade 3 form + bridges to Grades 2 and 4"),
        (4, "N456", "Distinct Grade 4 form + bridges to Grades 3 and 5"),
        (5, "N456", "Distinct Grade 5 form + bridges to Grades 4 and 6"),
        (6, "N456", "Distinct Grade 6 form + controlled link to Grade 5"),
    ], [0.7, 1.3, 4.8], 8.2, BLUE)
    add_bullets(doc, [
        "Build a grade-specific blueprint first; do not select items only from statistical ranking.",
        "Freeze bridge items exactly and document any translation, layout, administration, or scoring change.",
        "Balance the anchor subset across content domains, difficulty, response format, and independent stimuli.",
        "Pilot each assembled form by school and revalidate targeting, local dependence, item fit, and link sensitivity.",
    ])

    add_page_break(doc)
    doc.add_heading("9. Recommended work program", level=1)
    add_numbered(doc, [
        "Freeze anchors by link, not globally. Maintain separate pre/post, adjacent-grade, and trend anchor sets with version-lock evidence.",
        "Repair the two disconnected and fifteen unsupported/weak links before item freeze; retain eligible anchors and add only enough historical bank items to reach five pilot candidates per link.",
        "Create six grade-specific French operational forms with controlled bridges.",
        "Add stimulus_id, testlet_id, and rubric_id to the canonical item-version registry and preserve bundle structure.",
        "Have Ministry/content experts adjudicate the 101 proposed domain classifications and approve grade-specific blueprint weights.",
        "Verify every key/rubric and score reconstruction against source instruments and a machine-readable scoring manifest.",
        "Pilot shifted and revised items at the target grade by school; avoid random individual development/validation splits.",
        "Treat oral/fluency fields as separately scored components with explicit units, timing, caps, raters, and reliability.",
        "Use content-balanced anchor sets rather than the largest statistical set.",
        "Keep selection treatment-blind; run treatment-related DIF only after the primary instrument decisions are frozen.",
        "Maintain an exposure/change log and retire compromised anchors even if historical fit is strong.",
        "After form assembly, re-estimate reliability, dimensionality, local dependence, item fit, and link constants; publish sensitivity to alternate models and anchor sets.",
    ])
    add_callout(doc, "Approval gate.", "No linked score should be used for reporting until the source/version review, domain adjudication, scoring verification, and weak-link repair are signed off in the decision log.")

    add_page_break(doc)
    doc.add_heading("Appendix A. Methods and open issues", level=1)
    doc.add_heading("Methods", level=2)
    add_bullets(doc, [
        "Classical evidence: difficulty, nonresponse/don’t-know rates, corrected item-total association, alpha, and clustered uncertainty.",
        "Structural evidence: eigenvalue diagnostics, one-dimensional Rasch/2PL fit, adjusted-Q3 local dependence, and school-held-out stability.",
        "Anchor evidence: endpoint empirical quality, prompt-summary consistency, school-clustered logistic DIF on common-item rest scores, Holm-adjusted significance evidence, and mean–sigma sensitivity.",
        "Validation splits preserve school clustering. Seeds and software paths are recorded in tracked scripts and logs.",
    ])
    doc.add_heading("Open issues by severity", level=2)
    severity = Counter(r["severity"] for r in blocking)
    add_table(doc, ["Severity", "Count"], [(k.title(), v) for k, v in sorted(severity.items())], [3.4, 3.4], 8.4, BLUE)
    add_bullets(doc, [
        "Canonical source versions require collaborator confirmation despite file hashing.",
        "A machine-readable key/rubric manifest is not yet confirmed for every item-version.",
        "Stimulus/testlet/rubric metadata and exposure/change histories are incomplete.",
        "The inherited combined file and legacy code lineage remain deferred until the authorized legacy phase.",
        "Cross-year common-scale claims require Y1/Y2 response-data and full-version validation.",
    ])
    add_source_note(doc, "All detailed rows, provenance, warning evidence, sensitivity estimates, and Ministry edit columns are in the companion protected workbook.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def build_memo(work_root: Path, output_path: Path):
    outputs = work_root / "outputs/y3_ministry"
    score = load_json(outputs / "00_readiness/y3_score_reconstruction_summary.json")
    diag = load_json(outputs / "01_item_quality/y3_diagnostics_summary.json")
    anchor = load_json(outputs / "02_anchors/y3_anchor_summary.json")
    domain = load_json(outputs / "03_forms/y3_domain_completion_summary.json")
    ministry = load_json(outputs / "03_forms/y3_ministry_decision_summary.json")
    minimal = load_json(outputs / "03_forms/y3_minimal_change_summary.json")

    doc = Document()
    configure_document(doc, memo=True)
    add_masthead(doc, "Instrument Decisions for the Next Cycle", "Decision memo • Year 3 assessment audit • 4 August 2026", "To: Ministry assessment and curriculum teams")

    add_callout(
        doc,
        "Recommendation.",
        "Use a minimal-change pilot package: replace five clear problem items, hold six backups in reserve, retain current eligible anchors, and add only the historical items needed to repair links and tail coverage. Do not freeze a fully linked system until versions, keys, exposure, and pilot results are confirmed.",
        PALE_BLUE,
    )
    doc.add_heading("What is now available", level=1)
    add_bullets(doc, [
        f"A reproducible audit of {score['form_count']} forms and {score['binary_item_form_rows']:,} binary item×form contexts across baseline, pilot, and endline.",
        f"Item-level operational, classical, IRT, dimensionality, local-dependence, and school-split evidence; {diag['converged_model_rows']} of {diag['fitted_model_rows']} fitted model rows converged.",
        f"A completed proposed content/cognitive map for all {domain['item_form_grade_rows']} endline item×grade rows.",
        "Separate within-grade, adjacent-grade, and provisional trend anchor evidence with sensitivity results.",
        f"A Ministry decision workbook covering {ministry['endline_item_form_decision_rows']} endline rows, {ministry['one_grade_shift_rows']} one-grade-shift rows, and six proposed French forms.",
        "Three short decision sheets: Action Shortlist, Anchor Repairs, and Tail Coverage.",
    ])
    doc.add_heading("What the evidence says", level=1)
    add_key_value_table(doc, [
        ("Reliability", "0.827–0.972", "Strong, but high alpha does not prove one dimension"),
        ("Current anchor coverage", "16 / 33 adequate", "17 links need a targeted repair"),
        ("Minimal replacements", minimal["mandatory_replacement_rows"], f"Plus {minimal['conditional_backup_rows']} contingency backups"),
        ("Tail coverage", f"{minimal['tail_forms_currently_adequate']} / 18 adequate", f"{minimal['tail_forms_requiring_repair']} forms need targeted additions"),
        ("Local dependence", f"{diag['local_dependence_warning_pairs']:,} pairs", "Likely undocumented stories/images/oral tasks/rubrics"),
        ("Domain completion", f"{domain['item_form_grade_rows']} / {domain['item_form_grade_rows']}", f"{domain['confirmation_required_rows']} need Ministry confirmation"),
    ])
    doc.add_heading("Immediate Ministry decisions", level=1)
    add_numbered(doc, [
        "Confirm or revise the 101 proposed content/cognitive classifications and approve grade-specific blueprint weights.",
        f"Review the {minimal['mandatory_replacement_rows']} proposed replacements and keep the {minimal['conditional_backup_rows']} alternatives only as backups.",
        f"Approve targeted anchor repair for {minimal['deficient_links']} links: preserve eligible current anchors and add only enough unchanged bank items to reach five pilot candidates per link.",
        f"Approve tail additions only for the {minimal['tail_forms_requiring_repair']} forms that need them; the shortlist adds {minimal['net_new_tail_only_items']} tail-only items after reusing replacement and anchor candidates.",
        "Approve six distinct French grade forms with a controlled, unchanged bridge subset.",
        "Provide/confirm a machine-readable key and rubric manifest, including oral/fluency units, timing, raters, caps, and missing-value rules.",
    ])

    add_page_break(doc)
    doc.add_heading("Non-negotiable design rules", level=1)
    add_bullets(doc, [
        "An item is an anchor only for a specific link and only after exact version and exposure review.",
        "Do not edit the prompt, stimulus, options, key/rubric, layout, scoring, or administration of a frozen anchor without breaking and revalidating the link.",
        "Add stimulus_id/testlet_id/rubric_id and review passage/task bundles together.",
        "Use content-balanced anchors across difficulty and independent stimuli; do not choose the largest statistical set.",
        "Keep initial item selection treatment-blind. Treatment-related DIF comes later as a sensitivity analysis.",
        "Pilot by school at the target grade and re-estimate after the forms are assembled.",
    ])
    doc.add_heading("Minimal-change pilot package", level=1)
    add_table(doc, ["Decision", "Size", "Ministry action"], [
        ("Replace now", minimal["mandatory_replacement_rows"], "Choose one proposed bank replacement per row"),
        ("Conditional backup", minimal["conditional_backup_rows"], "Hold; use only if review confirms failure"),
        ("Anchor repair", f"{minimal['deficient_links']} links", "Approve the compact, link-specific candidate sets"),
        ("Tail-only additions", minimal["net_new_tail_only_items"], f"Target only the {minimal['tail_forms_requiring_repair']} deficient forms"),
        ("All other items", "Preserve", "Keep unless content, key, or version review finds a problem"),
    ], [2.0, 1.0, 3.8], 8.1, BLUE)
    doc.add_paragraph(
        "Historical performance creates a pilot shortlist, not automatic reuse. Exact prompt, stimulus, options, key/rubric, layout, scoring, administration, and exposure must match before any item is frozen as an anchor."
    )
    doc.add_heading("Instrument-development sequence", level=1)
    add_numbered(doc, [
        "Adjudicate domains and approve blueprints.",
        "Review the five replacement decisions and six backups.",
        "Repair deficient links only to the five-candidate pilot target.",
        "Verify keys/rubrics and assign testlet metadata.",
        "Assemble the six grade-specific forms and one-grade-shift baseline forms.",
        "Pilot by school at the target grade.",
        "Re-estimate item quality, structure, and link sensitivity; then freeze versions and anchors.",
    ])
    add_callout(doc, "Reporting boundary.", "Until these steps are complete, report form-specific scores and diagnostic evidence. Do not describe Year 1–Year 3 as a common scale merely because item IDs appear across years.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    if len(sys.argv) != 4:
        raise SystemExit("usage: build_documents.py <work_root> <technical_report.docx> <ministry_memo.docx>")
    work_root = Path(sys.argv[1])
    report_path = Path(sys.argv[2])
    memo_path = Path(sys.argv[3])
    build_report(work_root, report_path)
    build_memo(work_root, memo_path)
    print(json.dumps({"technical_report": str(report_path), "ministry_memo": str(memo_path)}))


if __name__ == "__main__":
    main()
