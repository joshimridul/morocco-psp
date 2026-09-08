#!/usr/bin/env python3
"""Build the short Ministry-facing note for the Year 3 pilot item sheet."""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
GRAY = "555555"
BLACK = "000000"
PALE_BLUE = "EAF2F8"


def set_run_font(run, *, size=None, color=BLACK, bold=None, italic=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, color, before, after in [
        ("Heading 1", 15, BLUE, 8, 4),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.05

    header = section.header.paragraphs[0]
    header.text = "PIONEER SCHOOLS | YEAR 3 ITEM REVIEW"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        set_run_font(run, size=8, color=GRAY, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("Ministry review note | Page ")
    set_run_font(run, size=8, color=GRAY)
    add_page_field(footer)


def add_title_block(doc):
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(0)
    label.paragraph_format.space_after = Pt(4)
    set_run_font(label.add_run("MINISTRY REVIEW NOTE"), size=8.5, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    set_run_font(
        title.add_run("Year 3 Pilot Items: What Is Ready and What We Need"),
        size=21,
        color=BLACK,
        bold=True,
    )

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)
    set_run_font(
        subtitle.add_run("A short guide to the accompanying two-tab item workbook"),
        size=12.5,
        color=GRAY,
        italic=True,
    )

    for label_text, value in [
        ("To", "Ministry assessment team"),
        ("Purpose", "Approve the pilot item set and complete the few remaining item details"),
        ("Date", "5 August 2026"),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        set_run_font(paragraph.add_run(f"{label_text}: "), size=10.5, bold=True)
        set_run_font(paragraph.add_run(value), size=10.5)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_numbered_step(doc, lead, detail):
    paragraph = doc.add_paragraph(style="List Number")
    set_run_font(paragraph.add_run(lead), bold=True)
    set_run_font(paragraph.add_run(detail))


def load_rows(work_root):
    source = (
        work_root
        / "outputs"
        / "y3_ministry"
        / "03_forms"
        / "y3_pilot_item_bank.csv"
    )
    with source.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def build(work_root, output_path):
    rows = load_rows(work_root)
    math_edits_path = work_root / "private_manifests" / "y3_math_surface_edits.csv"
    decisions_path = (
        work_root
        / "outputs"
        / "y3_ministry"
        / "03_forms"
        / "y3_ministry_item_decisions.csv"
    )
    domains_path = (
        work_root
        / "outputs"
        / "y3_ministry"
        / "03_forms"
        / "y3_completed_domain_map.csv"
    )
    with decisions_path.open(encoding="utf-8-sig", newline="") as handle:
        decisions = list(csv.DictReader(handle))
    with domains_path.open(encoding="utf-8-sig", newline="") as handle:
        domains = list(csv.DictReader(handle))
    with math_edits_path.open(encoding="utf-8-sig", newline="") as handle:
        math_edits = list(csv.DictReader(handle))
    if len(rows) != 42 or len(decisions) != 639 or len(domains) != 639 or len(math_edits) != 138:
        raise ValueError(
            f"Unexpected source counts: additions={len(rows)}, "
            f"decisions={len(decisions)}, domains={len(domains)}, math_edits={len(math_edits)}"
        )

    anchor_additions = sum("anchor:" in row["proposed_role"] for row in rows)
    immediate = sum("replacement now" in row["proposed_role"] for row in rows)
    reserve_replacements = sum(
        "conditional replacement" in row["proposed_role"]
        and "anchor:" not in row["proposed_role"]
        for row in rows
    )
    dual_role_anchors = sum(
        "conditional replacement" in row["proposed_role"]
        and "anchor:" in row["proposed_role"]
        for row in rows
    )
    tails = sum("tail support" in row["proposed_role"] for row in rows)
    unique_items = len({(row["subject"], row["item_id"]) for row in decisions})
    domain_confirmation = sum(
        row["requires_ministry_confirmation"] == "1" for row in domains
    )
    answer_needed = [
        row["item_id"]
        for row in rows
        if row["review_status"] == "answer/rubric confirmation required"
    ]
    item_needed = sum(
        row["review_status"] == "item and answer required" for row in rows
    )
    math_edit_ids = len({row["item_id"] for row in math_edits})
    math_visual_rows = sum(
        row["visual_or_layout_edit_required"].startswith("YES") for row in math_edits
    )

    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "Year 3 Pilot Items - Ministry Review Note"
    doc.core_properties.subject = "Instructions for the Ministry item review sheet"
    doc.core_properties.author = "Pioneer Schools assessment team"
    add_title_block(doc)

    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(7)
    lead.paragraph_format.space_after = Pt(7)
    lead.paragraph_format.left_indent = Inches(0.12)
    lead.paragraph_format.right_indent = Inches(0.12)
    shade_paragraph(lead, PALE_BLUE)
    set_run_font(lead.add_run("The short version: "), color=NAVY, bold=True)
    set_run_font(
        lead.add_run(
            f"the workbook has two tabs: the full {len(decisions)}-row Year 3 instrument "
            f"and {len(rows)} proposed additions. Complete only the yellow columns. "
            "All eligible non-anchor Maths revisions are now included; anchors remain unchanged."
        ),
        color=BLACK,
    )

    doc.add_heading("What we completed", level=1)
    add_bullet(
        doc,
        f"Reviewed all {len(decisions)} item-by-form/grade records, representing {unique_items} distinct Year 3 item IDs.",
    )
    add_bullet(
        doc,
        f"Completed content and cognitive domains for every row. {domain_confirmation} team-completed rows require Ministry confirmation.",
    )
    add_bullet(
        doc,
        f"Added a separate shortlist with {anchor_additions} historical anchor rows, {immediate} immediate replacements, {reserve_replacements} reserve replacement, and {tails} tail-support rows. {dual_role_anchors} active anchors can also replace a weak current item if needed.",
    )
    add_bullet(
        doc,
        "Marked every item that may serve as an anchor as frozen, and identified where surface changes are allowed.",
    )
    add_bullet(
        doc,
        f"For all {len(math_edits)} eligible non-anchor Maths rows ({math_edit_ids} item IDs), supplied the checked current key, proposed English revision, exact revised answer, and a plain-language description of the change. {math_visual_rows} rows explicitly require an artwork or layout update.",
    )

    doc.add_heading("What the Ministry needs to do", level=1)
    add_numbered_step(
        doc,
        "Review the full item bank. ",
        f"Confirm the {domain_confirmation} rows marked OUR COMPLETION - MINISTRY CONFIRM.",
    )
    add_numbered_step(
        doc,
        "Protect the anchors. ",
        "Do not change wording, names, objects, numbers, options, images, layout, answers, or scoring on rows marked ANCHOR SOMEWHERE - DO NOT EDIT.",
    )
    add_numbered_step(
        doc,
        "Review Proposed Additions. ",
        "Include and protect the active anchors, approve immediate replacements, and hold only the row marked RESERVE REPLACEMENT unless needed.",
    )
    add_numbered_step(
        doc,
        "Approve and implement the Maths revisions. ",
        f"Review our English wording and revised answer for the {len(math_edits)} eligible rows, translate the approved wording, and update the student and examiner artwork/layout where the sheet says so ({math_visual_rows} rows). Apply shared-stimulus edits as one bundle.",
    )
    add_numbered_step(
        doc,
        "Complete permitted language-item surface edits. ",
        "For Arabic and French non-anchors, change only a name, object, or number while preserving the skill, difficulty, options, and answer logic.",
    )
    add_numbered_step(
        doc,
        "Confirm keys and remaining French details. ",
        f"Resolve the rows marked TO VERIFY, confirm the answer or scoring rule for {', '.join(answer_needed)}, and provide one easy Grade 4 French item with its correct answer. Record decisions in the yellow columns.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return {
        "output": str(output_path),
        "full_rows": len(decisions),
        "addition_rows": len(rows),
        "anchor_addition_rows": anchor_additions,
        "reserve_replacement_rows": reserve_replacements,
        "dual_role_anchor_rows": dual_role_anchors,
        "domain_confirmation_rows": domain_confirmation,
        "answer_needed": len(answer_needed),
        "item_needed": item_needed,
        "math_edit_rows": len(math_edits),
        "math_edit_ids": math_edit_ids,
        "math_visual_rows": math_visual_rows,
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit(
            "Usage: build_simple_ministry_note.py <work_root> <output.docx>"
        )
    result = build(Path(sys.argv[1]), Path(sys.argv[2]))
    print(result)


if __name__ == "__main__":
    main()
