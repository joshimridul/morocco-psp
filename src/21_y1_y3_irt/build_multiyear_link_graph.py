#!/usr/bin/env python3
"""Build a version-aware Year 1--Year 3 IRT link graph.

Only hashes, item IDs, aggregate administration evidence, and IRT parameters are
written. Exact English-summary agreement is development evidence, not final
proof that prompts, stimuli, options, keys, rubrics, layout, scoring, and
exposure are invariant.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import heapq
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from itertools import combinations
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "00_inventory"))
from build_y3_registry import read_simple_paths  # noqa: E402
from path_guard import (  # noqa: E402
    assert_no_direct_identifier_columns,
    assert_output_path,
    assert_source_path,
)


SUBJECTS = {"Arabic": "arabic", "French": "french", "Maths": "maths"}
Y1_REFERENCE_TRANSFORM = {
    "Arabic": (-0.2400314007037489, 0.9127664340677355),
    "French": (-0.5166984276268118, 0.7985084593747850),
    "Maths": (-0.3713869649627327, 0.8408824092949523),
}
Y2_MAP_SPECS = (
    (
        "baseline",
        "Arabic",
        "3 - Data collection/01_Baseline/3 - Item map/Arabic/MCQ_Item_map_arabic_2024-11-14_ya.xlsx",
        "Arabic Baseline",
    ),
    (
        "baseline",
        "French",
        "3 - Data collection/01_Baseline/3 - Item map/French/MCQ_item map french 2024-10-11 ya.xlsx",
        "French Baseline",
    ),
    (
        "baseline",
        "Maths",
        "3 - Data collection/01_Baseline/3 - Item map/Math/MCQ_Item map maths 2024-10-11 ya.xlsx",
        "Maths Baseline",
    ),
    (
        "pilot",
        "",
        "3 - Data collection/03_Pilot/03_Item Maps/Item_maps_20250507_ks.xlsx",
        "Item Map",
    ),
    (
        "endline",
        "",
        "3 - Data collection/02_Endline/03_Item Maps/Item_maps_20250715_ks.xlsx",
        "Item Map",
    ),
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalized_header(value: Any) -> str:
    return re.sub(r"\s+", " ", clean(value).casefold()).strip()


def normalized_text(value: Any) -> str:
    return re.sub(r"\s+", " ", clean(value).casefold()).strip()


def text_hash(value: Any) -> str:
    value = normalized_text(value)
    return hashlib.sha256(value.encode("utf-8")).hexdigest() if value else ""


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: Iterable[dict[str, Any]], fields: list[str]) -> None:
    assert_no_direct_identifier_columns(fields)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def subject_from_item(item_id: str) -> str:
    prefix = item_id[:1].casefold()
    return {"a": "Arabic", "f": "French", "m": "Maths"}.get(prefix, "")


def read_compiled_hashes(path: Path) -> dict[tuple[str, str, int], set[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    result: dict[tuple[str, str, int], set[str]] = defaultdict(set)
    for values in iterator:
        row = dict(zip(headers, values))
        item_id = clean(row.get("LatestItemsID"))
        subject = clean(row.get("Subject"))
        year_text = clean(row.get("Year"))
        digest = text_hash(row.get("Thequestiontext"))
        match = re.fullmatch(r"Y([123])", year_text)
        if item_id and subject and match and digest:
            result[(item_id, subject, int(match.group(1)))].add(digest)
    workbook.close()
    return result


def read_y2_current_map_hashes(
    y2_root: Path, source_roots: list[Path]
) -> dict[tuple[str, str, str], set[str]]:
    result: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    id_headers = {
        "item_id_corrected",
        "new items id",
        "new item id",
        "items id",
        "new items id",
        "question id endline",
        "question id",
    }
    prompt_headers = {
        "the question text endline",
        "the question text",
        "question text",
    }
    subject_headers = {"the subject", "the subject ", "subject"}

    for wave, fixed_subject, relative_path, sheet_name in Y2_MAP_SPECS:
        path = assert_source_path(y2_root / relative_path, source_roots)
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing sheet {sheet_name!r} in {relative_path}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = list(next(iterator))
        headers = [normalized_header(value) for value in raw_headers]
        id_indexes = [index for index, value in enumerate(headers) if value in id_headers]
        prompt_indexes = [
            index for index, value in enumerate(headers) if value in prompt_headers
        ]
        subject_indexes = [
            index for index, value in enumerate(headers) if value in subject_headers
        ]
        if not id_indexes or not prompt_indexes:
            raise ValueError(f"Could not locate item ID and prompt columns in {relative_path}")
        for values in iterator:
            prompt = next(
                (clean(values[index]) for index in prompt_indexes if clean(values[index])),
                "",
            )
            digest = text_hash(prompt)
            if not digest:
                continue
            subject = fixed_subject or next(
                (
                    clean(values[index])
                    for index in subject_indexes
                    if clean(values[index])
                ),
                "",
            )
            item_ids = {clean(values[index]) for index in id_indexes if clean(values[index])}
            for item_id in item_ids:
                item_subject = subject or subject_from_item(item_id)
                if item_subject in SUBJECTS and re.match(r"^[afm][0-9]", item_id, re.I):
                    result[(wave, item_subject, item_id)].add(digest)
        workbook.close()
    return result


def exact_occurrence_hash(
    current_hashes: set[str], compiled_hashes: set[str]
) -> str:
    if len(current_hashes) == 1 and current_hashes == compiled_hashes:
        return next(iter(current_hashes))
    return ""


def version_id(subject: str, item_id: str, digest: str) -> str:
    suffix = digest[:16] if digest else "unresolved"
    return f"{subject}|{item_id}|{suffix}"


def selected_paths(
    nodes: set[str],
    node_subject: dict[str, str],
    direct_counts: dict[str, int],
    edge_counts: dict[tuple[str, str], int],
    mode: str,
) -> list[dict[str, Any]]:
    adjacency: dict[str, list[tuple[str, int]]] = defaultdict(list)
    for (left, right), count in edge_counts.items():
        if count >= 5:
            adjacency[left].append((right, count))
            adjacency[right].append((left, count))
    for values in adjacency.values():
        values.sort(key=lambda pair: (-pair[1], pair[0]))

    output: list[dict[str, Any]] = []
    for subject in SUBJECTS:
        subject_nodes = sorted(node for node in nodes if node_subject[node] == subject)
        root = f"Y1|reference|{subject}"
        # Cost favors the fewest links, then the largest path bottleneck, then a
        # deterministic lexical path. Direct links are always one edge.
        best: dict[str, tuple[int, int, str]] = {}
        parent: dict[str, str] = {}
        parent_n: dict[str, int] = {}
        root_n: dict[str, int] = {}
        bottleneck: dict[str, int] = {}
        path_text: dict[str, str] = {}
        queue: list[tuple[int, int, str, str]] = []
        for node in subject_nodes:
            count = direct_counts.get(node, 0)
            if count >= 5:
                path = f"{root}>{node}"
                cost = (1, -count, path)
                best[node] = cost
                parent[node] = root
                parent_n[node] = count
                root_n[node] = count
                bottleneck[node] = count
                path_text[node] = path
                heapq.heappush(queue, (*cost, node))

        while queue:
            depth, negative_bottleneck, path, node = heapq.heappop(queue)
            if best.get(node) != (depth, negative_bottleneck, path):
                continue
            for neighbor, count in adjacency.get(node, []):
                if node_subject.get(neighbor) != subject:
                    continue
                candidate_bottleneck = min(-negative_bottleneck, count)
                candidate_path = f"{path}>{neighbor}"
                candidate = (depth + 1, -candidate_bottleneck, candidate_path)
                if neighbor not in best or candidate < best[neighbor]:
                    best[neighbor] = candidate
                    parent[neighbor] = node
                    parent_n[neighbor] = count
                    root_n[neighbor] = root_n[node]
                    bottleneck[neighbor] = candidate_bottleneck
                    path_text[neighbor] = candidate_path
                    heapq.heappush(queue, (*candidate, neighbor))

        for node in subject_nodes:
            parts = node.split("|")
            reached = node in best
            output.append(
                {
                    "anchor_mode": mode,
                    "node_id": node,
                    "year": int(parts[0][1:]),
                    "wave": parts[1],
                    "subject": parts[2],
                    "administered_grade": parts[3],
                    "path_status": "linked_development_only" if reached else "no_path_with_5_anchors_per_edge",
                    "path_depth": best[node][0] if reached else "",
                    "parent_node_id": parent.get(node, ""),
                    "parent_edge_anchor_n": parent_n.get(node, ""),
                    "root_direct_anchor_n": root_n.get(node, ""),
                    "path_minimum_anchor_n": bottleneck.get(node, ""),
                    "selected_path": path_text.get(node, ""),
                    "full_version_review_status": "required",
                    "final_link_approved": 0,
                }
            )
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, required=True)
    args = parser.parse_args()
    paths = read_simple_paths(args.config.resolve())
    source_roots = [paths["y1_root"], paths["y2_root"], paths["y3_root"]]
    work_root = paths["work_root"]

    y2_registry_path = work_root / "derived/y1_y3_irt/y2_operational_item_registry.csv"
    y3_manifest_path = work_root / "outputs/y1_y3_irt/01_link_design/y3_y1_anchor_manifest.csv"
    y3_operational_path = work_root / "derived/y3_ministry/y3_item_operational_stats.csv"
    conflicts_path = work_root / "private_manifests/historical_bank_version_conflicts.csv"
    compiled_path = assert_source_path(
        paths["y3_root"]
        / "3 - Data collection/01_Baseline/3 - Item map/Y1-3 Item Map/Item_maps_compiled_16122025_eo.xlsx",
        source_roots,
    )
    for required in (y2_registry_path, y3_manifest_path, y3_operational_path):
        if not required.exists():
            raise FileNotFoundError(required)

    compiled = read_compiled_hashes(compiled_path)
    y2_current = read_y2_current_map_hashes(paths["y2_root"], source_roots)
    conflicts = {
        (row.get("subject", ""), row.get("candidate_id", ""))
        for row in read_csv(conflicts_path)
    }
    parameters: dict[tuple[str, str], tuple[float, float]] = {}
    for subject, slug in SUBJECTS.items():
        parameter_path = (
            work_root
            / f"outputs/y1_y3_irt/00_y1_reference/ster_probe/y1_{slug}_item_parameters.csv"
        )
        for row in read_csv(parameter_path):
            if row.get("fixed_in_combined_model") == "1":
                parameters[(subject, row["item_id"])] = (float(row["a"]), float(row["b"]))

    occurrences: list[dict[str, Any]] = []
    y2_rows = [
        row
        for row in read_csv(y2_registry_path)
        if row.get("variable_after_binary_scoring") == "1"
    ]
    y3_variable = {
        (row["wave"], row["subject"], row["grade"], row["item_id"])
        for row in read_csv(y3_operational_path)
        if row.get("item_type") == "binary"
        and float(row.get("n_correct") or 0) > 0
        and float(row.get("n_correct") or 0) < float(row.get("n_form") or 0)
    }
    y3_rows = [
        row
        for row in read_csv(y3_manifest_path)
        if (row["wave"], row["subject"], row["administered_grade"], row["item_id"])
        in y3_variable
    ]

    for row in y2_rows:
        subject = row["subject"]
        item_id = row["item_id"]
        year = 2
        current_hashes = y2_current.get((row["wave"], subject, item_id), set())
        year_hashes = compiled.get((item_id, subject, year), set())
        y1_hashes = compiled.get((item_id, subject, 1), set())
        exact_hash = exact_occurrence_hash(current_hashes, year_hashes)
        direct_hash = exact_hash if exact_hash and y1_hashes == {exact_hash} else ""
        conflict = int((subject, item_id) in conflicts)
        parameter = parameters.get((subject, item_id))
        direct_strict = int(parameter is not None and not conflict and bool(direct_hash))
        direct_provisional = int(parameter is not None and not conflict)
        a = parameter[0] if parameter else ""
        b = parameter[1] if parameter else ""
        ref_mean, ref_sd = Y1_REFERENCE_TRANSFORM[subject]
        occurrences.append(
            {
                "node_id": row["node_id"],
                "year": year,
                "wave": row["wave"],
                "subject": subject,
                "administered_grade": row["administered_grade"],
                "item_id": item_id,
                "strict_item_version_id": version_id(subject, item_id, exact_hash),
                "current_summary_hash": exact_hash,
                "current_summary_hash_count": len(current_hashes),
                "compiled_year_summary_hash_count": len(year_hashes),
                "compiled_y1_summary_hash_count": len(y1_hashes),
                "known_version_conflict": conflict,
                "direct_y1_strict_anchor": direct_strict,
                "direct_y1_provisional_anchor": direct_provisional,
                "y1_discrimination_a": a,
                "y1_difficulty_b": b,
                "mirt_intercept_d": -a * b if parameter else "",
                "y1_reference_mean": ref_mean,
                "y1_reference_sd": ref_sd,
                "full_version_review_status": "required",
                "final_anchor_approved": 0,
                "source_status": "provisional_development_input_requires_confirmation",
            }
        )

    for row in y3_rows:
        subject = row["subject"]
        item_id = row["item_id"]
        digest = row.get("y3_prompt_summary_sha256", "")
        conflict = int(row.get("known_version_conflict", "0") or 0)
        occurrences.append(
            {
                "node_id": f"Y3|{row['wave']}|{subject}|{row['administered_grade']}",
                "year": 3,
                "wave": row["wave"],
                "subject": subject,
                "administered_grade": row["administered_grade"],
                "item_id": item_id,
                "strict_item_version_id": version_id(subject, item_id, digest),
                "current_summary_hash": digest,
                "current_summary_hash_count": int(bool(digest)),
                "compiled_year_summary_hash_count": int(row.get("y3_compiled_summary_hash_count", "0") or 0),
                "compiled_y1_summary_hash_count": int(row.get("y1_compiled_summary_hash_count", "0") or 0),
                "known_version_conflict": conflict,
                "direct_y1_strict_anchor": int(row["strict_development_anchor"]),
                "direct_y1_provisional_anchor": int(row["provisional_id_anchor"]),
                "y1_discrimination_a": row["y1_discrimination_a"],
                "y1_difficulty_b": row["y1_difficulty_b"],
                "mirt_intercept_d": row["mirt_intercept_d"],
                "y1_reference_mean": row["y1_reference_mean"],
                "y1_reference_sd": row["y1_reference_sd"],
                "full_version_review_status": "required",
                "final_anchor_approved": 0,
                "source_status": "provisional_development_input_requires_confirmation",
            }
        )

    occurrences.sort(
        key=lambda row: (
            row["subject"],
            row["year"],
            row["wave"],
            int(row["administered_grade"]),
            row["item_id"],
        )
    )
    seen_occurrences = set()
    unique_occurrences = []
    for row in occurrences:
        key = (row["node_id"], row["item_id"])
        if key not in seen_occurrences:
            seen_occurrences.add(key)
            unique_occurrences.append(row)
    occurrences = unique_occurrences

    by_subject_item: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in occurrences:
        by_subject_item[(row["subject"], row["item_id"])].append(row)

    link_items: list[dict[str, Any]] = []
    for (subject, item_id), rows in by_subject_item.items():
        if len(rows) < 2:
            continue
        for left, right in combinations(sorted(rows, key=lambda row: row["node_id"]), 2):
            if left["node_id"] == right["node_id"]:
                continue
            strict = int(
                not left["known_version_conflict"]
                and not right["known_version_conflict"]
                and bool(left["current_summary_hash"])
                and left["current_summary_hash"] == right["current_summary_hash"]
                and left["compiled_year_summary_hash_count"] == 1
                and right["compiled_year_summary_hash_count"] == 1
            )
            provisional = int(
                not left["known_version_conflict"] and not right["known_version_conflict"]
            )
            if not provisional:
                continue
            node_a, node_b = sorted((left["node_id"], right["node_id"]))
            link_items.append(
                {
                    "subject": subject,
                    "node_a": node_a,
                    "node_b": node_b,
                    "item_id": item_id,
                    "strict_item_version_id": left["strict_item_version_id"] if strict else "",
                    "strict_summary_link": strict,
                    "provisional_id_link": provisional,
                    "known_version_conflict": 0,
                    "full_version_review_status": "required",
                    "final_link_item_approved": 0,
                }
            )

    edge_groups: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in link_items:
        edge_groups[(row["subject"], row["node_a"], row["node_b"])].append(row)
    edges = []
    for (subject, node_a, node_b), rows in sorted(edge_groups.items()):
        edges.append(
            {
                "subject": subject,
                "node_a": node_a,
                "node_b": node_b,
                "strict_summary_anchor_n": sum(row["strict_summary_link"] for row in rows),
                "provisional_id_anchor_n": sum(row["provisional_id_link"] for row in rows),
                "strict_edge_eligible": int(sum(row["strict_summary_link"] for row in rows) >= 5),
                "provisional_edge_eligible": int(sum(row["provisional_id_link"] for row in rows) >= 5),
                "full_version_review_status": "required",
                "final_edge_approved": 0,
            }
        )

    nodes = {row["node_id"] for row in occurrences}
    node_subject = {row["node_id"]: row["subject"] for row in occurrences}
    direct_strict = Counter()
    direct_provisional = Counter()
    for row in occurrences:
        direct_strict[row["node_id"]] += int(row["direct_y1_strict_anchor"])
        direct_provisional[row["node_id"]] += int(row["direct_y1_provisional_anchor"])
    strict_edges = {
        (row["node_a"], row["node_b"]): int(row["strict_summary_anchor_n"])
        for row in edges
    }
    provisional_edges = {
        (row["node_a"], row["node_b"]): int(row["provisional_id_anchor_n"])
        for row in edges
    }
    paths_rows = selected_paths(
        nodes, node_subject, dict(direct_strict), strict_edges, "strict_summary"
    ) + selected_paths(
        nodes,
        node_subject,
        dict(direct_provisional),
        provisional_edges,
        "provisional_id",
    )
    paths_rows.sort(
        key=lambda row: (
            row["anchor_mode"],
            row["subject"],
            row["year"],
            row["wave"],
            int(row["administered_grade"]),
        )
    )

    out_dir = assert_output_path(
        work_root / "outputs/y1_y3_irt/01_multiyear_link_graph",
        work_root,
        source_roots,
    )
    outputs = {
        "occurrences": out_dir / "multiyear_item_occurrence_manifest.csv",
        "link_items": out_dir / "multiyear_link_item_manifest.csv",
        "edges": out_dir / "multiyear_node_edges.csv",
        "paths": out_dir / "multiyear_selected_paths.csv",
        "review": out_dir / "multiyear_version_review_queue.csv",
        "selected_review": out_dir / "selected_path_version_review_queue.csv",
        "summary": out_dir / "multiyear_link_graph_summary.json",
    }
    write_csv(outputs["occurrences"], occurrences, list(occurrences[0]))
    write_csv(outputs["link_items"], link_items, list(link_items[0]))
    write_csv(outputs["edges"], edges, list(edges[0]))
    write_csv(outputs["paths"], paths_rows, list(paths_rows[0]))

    review_rows = []
    for row in occurrences:
        if row["direct_y1_strict_anchor"] or row["direct_y1_provisional_anchor"]:
            review_rows.append(
                {
                    "review_type": "direct_y1_anchor",
                    "subject": row["subject"],
                    "item_id": row["item_id"],
                    "node_a": f"Y1|reference|{row['subject']}",
                    "node_b": row["node_id"],
                    "automated_prompt_evidence": "exact_single_summary" if row["direct_y1_strict_anchor"] else "id_only",
                    "prompt_review": "required",
                    "stimulus_review": "required",
                    "options_review": "required",
                    "key_or_rubric_review": "required",
                    "administration_scoring_layout_review": "required",
                    "exposure_review": "required",
                    "final_decision": "not_approved",
                }
            )
    for row in link_items:
        if row["strict_summary_link"]:
            review_rows.append(
                {
                    "review_type": "bridge_anchor",
                    "subject": row["subject"],
                    "item_id": row["item_id"],
                    "node_a": row["node_a"],
                    "node_b": row["node_b"],
                    "automated_prompt_evidence": "exact_single_summary",
                    "prompt_review": "required",
                    "stimulus_review": "required",
                    "options_review": "required",
                    "key_or_rubric_review": "required",
                    "administration_scoring_layout_review": "required",
                    "exposure_review": "required",
                    "final_decision": "not_approved",
                }
            )
    review_rows.sort(key=lambda row: (row["subject"], row["item_id"], row["node_a"], row["node_b"]))
    write_csv(outputs["review"], review_rows, list(review_rows[0]))

    selected_review_rows: list[dict[str, Any]] = []
    for mode in ("strict_summary", "provisional_id"):
        direct_flag = (
            "direct_y1_strict_anchor"
            if mode == "strict_summary"
            else "direct_y1_provisional_anchor"
        )
        bridge_flag = (
            "strict_summary_link"
            if mode == "strict_summary"
            else "provisional_id_link"
        )
        targets = [
            row
            for row in paths_rows
            if row["anchor_mode"] == mode
            and row["year"] == 3
            and row["wave"] in {"baseline", "endline"}
            and row["path_status"] == "linked_development_only"
        ]
        active_nodes = {
            component
            for row in targets
            for component in row["selected_path"].split(">")
            if component.startswith(("Y2|", "Y3|"))
        }
        selected_by_node = {
            row["node_id"]: row
            for row in paths_rows
            if row["anchor_mode"] == mode and row["node_id"] in active_nodes
        }
        for node_id in sorted(active_nodes):
            path_row = selected_by_node[node_id]
            parent = path_row["parent_node_id"]
            if parent.startswith("Y1|reference|"):
                items = [
                    row
                    for row in occurrences
                    if row["node_id"] == node_id and row[direct_flag] == 1
                ]
                for item in items:
                    selected_review_rows.append(
                        {
                            "anchor_mode": mode,
                            "path_depth": path_row["path_depth"],
                            "review_type": "selected_direct_y1_anchor",
                            "subject": item["subject"],
                            "item_id": item["item_id"],
                            "node_a": parent,
                            "node_b": node_id,
                            "automated_prompt_evidence": (
                                "exact_single_summary"
                                if mode == "strict_summary"
                                else "id_only"
                            ),
                            "prompt_review": "required",
                            "stimulus_review": "required",
                            "options_review": "required",
                            "key_or_rubric_review": "required",
                            "administration_scoring_layout_review": "required",
                            "exposure_review": "required",
                            "final_decision": "not_approved",
                        }
                    )
            else:
                node_a, node_b = sorted((parent, node_id))
                items = [
                    row
                    for row in link_items
                    if row["node_a"] == node_a
                    and row["node_b"] == node_b
                    and row[bridge_flag] == 1
                ]
                for item in items:
                    selected_review_rows.append(
                        {
                            "anchor_mode": mode,
                            "path_depth": path_row["path_depth"],
                            "review_type": "selected_bridge_anchor",
                            "subject": item["subject"],
                            "item_id": item["item_id"],
                            "node_a": parent,
                            "node_b": node_id,
                            "automated_prompt_evidence": (
                                "exact_single_summary"
                                if mode == "strict_summary"
                                else "id_only"
                            ),
                            "prompt_review": "required",
                            "stimulus_review": "required",
                            "options_review": "required",
                            "key_or_rubric_review": "required",
                            "administration_scoring_layout_review": "required",
                            "exposure_review": "required",
                            "final_decision": "not_approved",
                        }
                    )
    selected_review_rows.sort(
        key=lambda row: (
            row["anchor_mode"],
            row["subject"],
            int(row["path_depth"]),
            row["node_b"],
            row["item_id"],
        )
    )
    write_csv(
        outputs["selected_review"],
        selected_review_rows,
        list(selected_review_rows[0]),
    )

    summary = {
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "node_n": len(nodes),
        "node_item_occurrence_n": len(occurrences),
        "link_item_pair_n": len(link_items),
        "edge_n": len(edges),
        "version_review_row_n": len(review_rows),
        "selected_path_version_review_row_n": len(selected_review_rows),
        "direct_strict_node_n": sum(value >= 5 for value in direct_strict.values()),
        "direct_provisional_node_n": sum(value >= 5 for value in direct_provisional.values()),
        "strict_linked_node_n": sum(
            row["path_status"] == "linked_development_only"
            for row in paths_rows
            if row["anchor_mode"] == "strict_summary"
        ),
        "provisional_linked_node_n": sum(
            row["path_status"] == "linked_development_only"
            for row in paths_rows
            if row["anchor_mode"] == "provisional_id"
        ),
        "strict_linked_y3_baseline_endline_node_n": sum(
            row["path_status"] == "linked_development_only"
            and row["year"] == 3
            and row["wave"] in {"baseline", "endline"}
            for row in paths_rows
            if row["anchor_mode"] == "strict_summary"
        ),
        "provisional_linked_y3_baseline_endline_node_n": sum(
            row["path_status"] == "linked_development_only"
            and row["year"] == 3
            and row["wave"] in {"baseline", "endline"}
            for row in paths_rows
            if row["anchor_mode"] == "provisional_id"
        ),
        "status": "development_only_full_version_review_required",
        "interpretation": (
            "Every selected edge has at least five automated development anchors. "
            "No edge or anchor is final until prompt, stimulus, options, key/rubric, "
            "administration, scoring, layout, and exposure are reviewed."
        ),
    }
    outputs["summary"].write_text(
        json.dumps(summary, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary))


if __name__ == "__main__":
    main()
