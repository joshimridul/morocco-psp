#!/usr/bin/env python3
"""Build the direct Year 1 fixed-parameter to Year 3 anchor manifest.

The output contains item IDs, hashes, and model parameters but no item text or
student data. Matching IDs and English summaries support a development link;
they never establish final operational equivalence by themselves.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "00_inventory"))
from build_y3_registry import read_simple_paths  # noqa: E402
from path_guard import (  # noqa: E402
    assert_no_direct_identifier_columns,
    assert_output_path,
    assert_source_path,
)


SUBJECTS = {
    "Arabic": "arabic",
    "French": "french",
    "Maths": "maths",
}

Y1_REFERENCE_TRANSFORM = {
    "Arabic": (-0.2400314007037489, 0.9127664340677355),
    "French": (-0.5166984276268118, 0.7985084593747850),
    "Maths": (-0.3713869649627327, 0.8408824092949523),
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalized_text(value: Any) -> str:
    return re.sub(r"\s+", " ", clean(value).casefold())


def text_hash(value: Any) -> str:
    normalized = normalized_text(value)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest() if normalized else ""


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: Iterable[dict[str, Any]], fields: list[str]) -> None:
    assert_no_direct_identifier_columns(fields)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read_compiled_summary_hashes(path: Path) -> dict[tuple[str, str, str], set[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    result: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for values in iterator:
        row = dict(zip(headers, values))
        item_id = clean(row.get("LatestItemsID"))
        subject = clean(row.get("Subject"))
        year = clean(row.get("Year"))
        digest = text_hash(row.get("Thequestiontext"))
        if item_id and subject and year and digest:
            result[(item_id, subject, year)].add(digest)
    workbook.close()
    return result


def summary_equivalence_status(
    y1_hashes: set[str], y3_hashes: set[str], current_hash: str
) -> str:
    """Classify map-summary evidence without claiming full item equivalence."""
    if (
        len(y1_hashes) == 1
        and y1_hashes == y3_hashes
        and current_hash
        and current_hash in y1_hashes
    ):
        return "exact_single_summary"
    overlap = y1_hashes & y3_hashes
    if overlap and current_hash in overlap:
        return "overlapping_summary_set_requires_version_resolution"
    return "summary_mismatch_or_missing"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--github-root", type=Path, required=True)
    args = parser.parse_args()

    paths = read_simple_paths(args.config.resolve())
    github_root = args.github_root.resolve()
    source_roots = [paths["y1_root"], paths["y2_root"], paths["y3_root"]]
    work_root = paths["work_root"]

    registry_path = work_root / "derived/y3_ministry/y3_item_version_registry.csv"
    operational_path = work_root / "derived/y3_ministry/y3_item_operational_stats.csv"
    compiled_path = (
        paths["y3_root"]
        / "3 - Data collection/01_Baseline/3 - Item map/Y1-3 Item Map/Item_maps_compiled_16122025_eo.xlsx"
    )
    conflicts_path = work_root / "private_manifests/historical_bank_version_conflicts.csv"
    assert_source_path(compiled_path, source_roots)
    for required in (registry_path, operational_path):
        if not required.exists():
            raise FileNotFoundError(required)

    summary_hashes = read_compiled_summary_hashes(compiled_path)
    registry = read_csv_rows(registry_path)
    operational = read_csv_rows(operational_path)
    binary_keys = {
        (row["wave"], row["subject"], row["grade"], row["item_id"])
        for row in operational
        if row.get("item_type") == "binary"
    }
    known_conflicts = set()
    if conflicts_path.exists():
        known_conflicts = {
            (row.get("subject", ""), row.get("candidate_id", ""))
            for row in read_csv_rows(conflicts_path)
        }

    parameters: dict[tuple[str, str], dict[str, str]] = {}
    for subject, subject_slug in SUBJECTS.items():
        parameter_path = (
            work_root
            / f"outputs/y1_y3_irt/00_y1_reference/ster_probe/y1_{subject_slug}_item_parameters.csv"
        )
        if not parameter_path.exists():
            raise FileNotFoundError(parameter_path)
        for row in read_csv_rows(parameter_path):
            if row.get("fixed_in_combined_model") == "1":
                parameters[(subject, row["item_id"])] = row

    manifest: list[dict[str, Any]] = []
    seen = set()
    for row in registry:
        key = (row["wave"], row["subject"], row["form_grade"], row["item_id"])
        if key in seen or key not in binary_keys:
            continue
        seen.add(key)
        subject = row["subject"]
        item_id = row["item_id"]
        parameter = parameters.get((subject, item_id))
        y1_hashes = summary_hashes.get((item_id, subject, "Y1"), set())
        y3_hashes = summary_hashes.get((item_id, subject, "Y3"), set())
        status = summary_equivalence_status(
            y1_hashes, y3_hashes, row.get("prompt_summary_sha256", "")
        )
        conflict = int((subject, item_id) in known_conflicts)
        fixed_match = int(parameter is not None)
        strict = int(fixed_match == 1 and conflict == 0 and status == "exact_single_summary")
        provisional = int(fixed_match == 1 and conflict == 0)
        ref_mean, ref_sd = Y1_REFERENCE_TRANSFORM[subject]
        a = float(parameter["a"]) if parameter else None
        b = float(parameter["b"]) if parameter else None
        manifest.append(
            {
                "wave": row["wave"],
                "subject": subject,
                "administered_grade": row["form_grade"],
                "form_group": row.get("form_group", ""),
                "item_id": item_id,
                "item_type": "binary",
                "y1_fixed_parameter_match": fixed_match,
                "y1_discrimination_a": a if a is not None else "",
                "y1_difficulty_b": b if b is not None else "",
                "mirt_intercept_d": (-a * b) if a is not None and b is not None else "",
                "y1_reference_mean": ref_mean,
                "y1_reference_sd": ref_sd,
                "summary_equivalence_status": status,
                "known_version_conflict": conflict,
                "strict_development_anchor": strict,
                "provisional_id_anchor": provisional,
                "full_version_review_status": "required",
                "final_anchor_approved": 0,
                "scale_claim_status": "development_only_not_final_common_scale",
                "y3_prompt_summary_sha256": row.get("prompt_summary_sha256", ""),
                "y1_compiled_summary_hash_count": len(y1_hashes),
                "y3_compiled_summary_hash_count": len(y3_hashes),
            }
        )

    manifest.sort(
        key=lambda row: (
            row["subject"],
            row["wave"],
            int(row["administered_grade"]),
            row["item_id"],
        )
    )
    manifest_fields = list(manifest[0])
    out_dir = assert_output_path(
        work_root / "outputs/y1_y3_irt/01_link_design", work_root, source_roots
    )
    manifest_path = assert_output_path(
        out_dir / "y3_y1_anchor_manifest.csv", work_root, source_roots
    )
    write_csv(manifest_path, manifest, manifest_fields)

    by_node: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in manifest:
        by_node[(row["wave"], row["subject"], row["administered_grade"])].append(row)
    node_rows: list[dict[str, Any]] = []
    for (wave, subject, grade), rows in sorted(
        by_node.items(), key=lambda pair: (pair[0][1], pair[0][0], int(pair[0][2]))
    ):
        fixed = [row for row in rows if row["y1_fixed_parameter_match"] == 1]
        strict = [row for row in rows if row["strict_development_anchor"] == 1]
        provisional = [row for row in rows if row["provisional_id_anchor"] == 1]
        difficulties = [float(row["y1_difficulty_b"]) for row in strict]
        node_rows.append(
            {
                "wave": wave,
                "subject": subject,
                "administered_grade": grade,
                "binary_item_n": len(rows),
                "direct_y1_fixed_id_n": len(fixed),
                "strict_development_anchor_n": len(strict),
                "provisional_id_anchor_n": len(provisional),
                "strict_anchor_difficulty_min": min(difficulties) if difficulties else "",
                "strict_anchor_difficulty_max": max(difficulties) if difficulties else "",
                "strict_anchor_difficulty_range": (
                    max(difficulties) - min(difficulties) if len(difficulties) >= 2 else ""
                ),
                "direct_strict_link_status": (
                    "adequate_count_pending_full_version_review"
                    if len(strict) >= 5
                    else "fewer_than_5_strict_summary_matches"
                ),
                "final_scale_status": "not_established",
            }
        )

    node_path = assert_output_path(
        out_dir / "y3_y1_node_link_summary.csv", work_root, source_roots
    )
    write_csv(node_path, node_rows, list(node_rows[0]))

    summary = {
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "manifest_rows": len(manifest),
        "node_rows": len(node_rows),
        "fixed_id_match_rows": sum(row["y1_fixed_parameter_match"] for row in manifest),
        "strict_development_anchor_rows": sum(
            row["strict_development_anchor"] for row in manifest
        ),
        "provisional_id_anchor_rows": sum(row["provisional_id_anchor"] for row in manifest),
        "summary_status_counts": Counter(
            row["summary_equivalence_status"] for row in manifest
        ),
        "nodes_with_at_least_5_strict_development_anchors": sum(
            row["strict_development_anchor_n"] >= 5 for row in node_rows
        ),
        "nodes_with_at_least_5_provisional_id_anchors": sum(
            row["provisional_id_anchor_n"] >= 5 for row in node_rows
        ),
        "interpretation": (
            "Strict development anchors require a fixed Year 1 parameter, one matching "
            "compiled English summary in Y1 and Y3, agreement with the current Y3 map, "
            "and no known version conflict. Prompt, stimulus, options, key or rubric, "
            "layout, administration, scoring, and exposure still require review."
        ),
    }
    summary_path = assert_output_path(
        out_dir / "y3_y1_link_design_summary.json", work_root, source_roots
    )
    summary_path.write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, default=dict) + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "status": "complete",
                "manifest": str(manifest_path),
                "node_summary": str(node_path),
                **summary,
            },
            ensure_ascii=False,
            default=dict,
        )
    )


if __name__ == "__main__":
    main()
